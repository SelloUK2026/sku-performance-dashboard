from __future__ import annotations

import base64
import io
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import (  # noqa: E402
    NON_EXISTING_SKU,
    aggregate_lifetime_metrics,
    align_channeladvisor_prices,
    ca_price_sku,
    calculate_candidate,
    commission_rows_from_workbook,
    map_imported_rows,
    normalise_date,
    resolve_wooper_sku,
    select_suggested_freight,
    valid_basic_authorization,
)


CRITERIA = {
    "min_grade": 0,
    "min_stock": 0,
    "min_months": 0,
    "min_discount": 0,
    "max_return_rate": 1,
    "return_volume_threshold": 20,
    "margin_volume_threshold": 10,
    "categories": [],
    "max_discount": 0.25,
    "default_commission": 0.264,
    "rounding": True,
    "wms_rate": 0.08,
    "vat_rate": 0.20,
    "grade_margins": {
        "0": 0.12,
        "1": 0.25,
        "2": 0.20,
        "3": 0.15,
        "4": 0.12,
        "5": 0.08,
        "6": 0.05,
        "7": 0.03,
    },
}


CORRECTED_WMS_RESULTS = {
    "AI1005-BK-UK": (0.25, 22.50, 0.12123243242666658),
    "AI1010-WH-UK": (0.03665972461900091, 74.15, 0.029820412763317653),
    "AI1012-WH-UK": (0.031055767296581993, 29.05, 0.1998251678898452),
    "AP0029-BK-UK": (0.15618844144771016, 46.35, 0.11926213592233004),
    "AP0029-RD-UK": (0.12434048679391552, 46.35, 0.11926213592233004),
    "AP0034-BK-UK": (0.24554427127518708, 33.85, 0.11816545225406203),
    "AP0034-RD-UK": (0.003324188164490094, 40.85, 0.24995072358384335),
    "AP0043-UK": (0.2212771039956496, 38.85, 0.11865269349806949),
    "AP0044-UK": (0.21725660389947632, 41.45, 0.029495778045838535),
    "BATH1010-WH-UK": (-0.0021185766505094072, 27.95, 0.11809573371019681),
}


class PromotionCalculationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows = json.loads((ROOT / "sample_data.json").read_text(encoding="utf-8"))["rows"]

    def test_basic_authorization_requires_exact_credentials(self):
        valid_header = "Basic " + base64.b64encode(b"sello:secret").decode("ascii")
        invalid_header = "Basic " + base64.b64encode(b"sello:wrong").decode("ascii")
        with (
            patch("app.PROMOTION_AUTH_USERNAME", "sello"),
            patch("app.PROMOTION_AUTH_PASSWORD", "secret"),
        ):
            self.assertTrue(valid_basic_authorization(valid_header))
            self.assertFalse(valid_basic_authorization(invalid_header))
            self.assertFalse(valid_basic_authorization("Bearer token"))

    def test_sample_rows_match_corrected_wms_results(self):
        for row in self.rows:
            with self.subTest(sku=row["sku"]):
                expected_discount, expected_price, expected_margin = (
                    CORRECTED_WMS_RESULTS[row["sku"]]
                )
                actual = calculate_candidate(row, CRITERIA)
                self.assertAlmostEqual(actual["suggested_discount"], expected_discount, places=10)
                self.assertAlmostEqual(actual["promo_price"], expected_price, places=8)
                self.assertAlmostEqual(actual["promo_margin"], expected_margin, places=10)

    def test_override_discount_recalculates_price_and_margin(self):
        row = dict(self.rows[1], override_discount=0.10)
        actual = calculate_candidate(row, CRITERIA)
        self.assertAlmostEqual(actual["final_discount"], 0.10)
        self.assertAlmostEqual(actual["promo_price"], 69.25)
        self.assertLess(actual["promo_margin"], actual["target_margin"])

    def test_vat_excluded_import_normalises_to_the_same_calculation_price(self):
        included = calculate_candidate(self.rows[1], CRITERIA)
        excluded_row = dict(self.rows[1], price=self.rows[1]["price"] / 1.20)
        excluded_criteria = dict(
            CRITERIA,
            input_price_includes_vat=False,
            export_price_includes_vat=True,
        )
        excluded = calculate_candidate(excluded_row, excluded_criteria)

        self.assertAlmostEqual(
            excluded["price_including_vat"],
            self.rows[1]["price"],
            places=8,
        )
        self.assertAlmostEqual(
            excluded["suggested_discount"],
            included["suggested_discount"],
            places=10,
        )
        self.assertAlmostEqual(excluded["promo_price"], included["promo_price"], places=8)
        self.assertAlmostEqual(excluded["promo_margin"], included["promo_margin"], places=10)

    def test_vat_excluded_export_converts_only_the_output_price(self):
        included = calculate_candidate(self.rows[1], CRITERIA)
        excluded = calculate_candidate(
            self.rows[1],
            dict(CRITERIA, export_price_includes_vat=False),
        )

        self.assertAlmostEqual(
            excluded["promo_price"],
            included["promo_price"] / 1.20,
            places=8,
        )
        self.assertAlmostEqual(
            excluded["suggested_discount"],
            included["suggested_discount"],
            places=10,
        )
        self.assertAlmostEqual(excluded["promo_profit"], included["promo_profit"], places=10)
        self.assertAlmostEqual(excluded["promo_margin"], included["promo_margin"], places=10)

    def test_user_can_choose_ca_price_or_offer_price(self):
        row = dict(self.rows[1], offer_price=60, ca_price=90)
        ca_result = calculate_candidate(row, dict(CRITERIA, price_source="ca"))
        offer_result = calculate_candidate(row, dict(CRITERIA, price_source="offer"))

        self.assertEqual(ca_result["calculation_price_source"], "ca")
        self.assertEqual(ca_result["price_including_vat"], 90)
        self.assertEqual(offer_result["calculation_price_source"], "offer")
        self.assertEqual(offer_result["price_including_vat"], 60)
        self.assertNotEqual(
            ca_result["suggested_discount"],
            offer_result["suggested_discount"],
        )

    def test_ca_price_is_the_vat_included_default(self):
        row = dict(self.rows[1], offer_price=60, ca_price=90)
        result = calculate_candidate(
            row,
            dict(CRITERIA, input_price_includes_vat=False),
        )

        self.assertEqual(result["requested_price_source"], "ca")
        self.assertEqual(result["calculation_price_source"], "ca")
        self.assertEqual(result["price_including_vat"], 90)

    def test_missing_offer_price_falls_back_to_ca_price(self):
        row = dict(self.rows[1], price=None, offer_price=None, ca_price=76.99)
        result = calculate_candidate(row, dict(CRITERIA, price_source="offer"))

        self.assertEqual(result["calculation_price_source"], "ca")
        self.assertEqual(result["price_including_vat"], 76.99)
        self.assertIn("Offer price unavailable; CA price used", result["warnings"])

    def test_missing_offer_and_ca_prices_do_not_claim_a_fallback(self):
        row = dict(self.rows[1], price=None, offer_price=None, ca_price=None)
        result = calculate_candidate(row, dict(CRITERIA, price_source="offer"))

        self.assertIsNone(result["calculation_price_source"])
        self.assertNotIn("Offer price unavailable; CA price used", result["warnings"])
        self.assertIn(
            "Price or margin settings make reverse pricing impossible",
            result["reasons"],
        )

    def test_max_discount_uses_ten_pence_rounding(self):
        actual = calculate_candidate(self.rows[0], CRITERIA)
        self.assertEqual(actual["final_discount"], 0.25)
        self.assertAlmostEqual(actual["promo_price"], 22.50)

    def test_nomination_filters_are_separate_from_pricing(self):
        criteria = dict(CRITERIA, min_stock=1, min_grade=4, min_discount=0.05)
        excluded = calculate_candidate(self.rows[2], criteria)
        eligible = calculate_candidate(self.rows[7], criteria)
        self.assertFalse(excluded["eligible"])
        self.assertIn("Grade below threshold", excluded["reasons"])
        self.assertTrue(eligible["eligible"])

    def test_multiselect_filters_match_any_value_within_each_filter(self):
        row = dict(
            self.rows[0],
            main_category="Home Appliances",
            subcategory="Dehumidifiers",
            brand="Spector",
        )
        included = calculate_candidate(
            row,
            dict(
                CRITERIA,
                main_categories=["Bathroom", "Home Appliances"],
                subcategories=["Juicers", "Dehumidifiers"],
                brands=["Levede", "Spector"],
            ),
        )
        excluded = calculate_candidate(
            row,
            dict(
                CRITERIA,
                main_categories=["Home Appliances"],
                subcategories=["Dehumidifiers"],
                brands=["Levede"],
            ),
        )

        self.assertTrue(included["eligible"])
        self.assertFalse(excluded["eligible"])
        self.assertIn("Outside selected brand", excluded["reasons"])

    def test_single_value_filter_requests_remain_supported(self):
        row = dict(
            self.rows[0],
            main_category="Home Appliances",
            subcategory="Dehumidifiers",
            brand="Spector",
        )
        actual = calculate_candidate(
            row,
            dict(
                CRITERIA,
                main_category="Home Appliances",
                subcategory="Dehumidifiers",
                brand="Spector",
            ),
        )

        self.assertTrue(actual["eligible"])

    def test_return_rate_at_threshold_is_a_review_warning(self):
        row = dict(self.rows[1], return_rate=0.06, sold_qty=0)
        actual = calculate_candidate(
            row,
            dict(CRITERIA, max_return_rate=0.06),
        )

        self.assertTrue(actual["eligible"])
        self.assertTrue(actual["return_rate_review"])
        self.assertIn(
            "Return rate is at or above 6% review threshold",
            actual["warnings"],
        )
        self.assertNotIn("High return rate on meaningful volume", actual["reasons"])

    def test_lifetime_margin_gap_over_five_points_is_a_review_warning(self):
        row = {
            "sku": "CD1024-DGY-66X72-UK",
            "grade": 4,
            "cogs": 9.69,
            "avg_freight": 1.71,
            "ca_price": 25.99,
            "return_rate": 0,
            "lifetime_profit_margin": 0.171,
        }
        result = calculate_candidate(
            row,
            dict(
                CRITERIA,
                default_commission=0.24,
                price_source="ca",
                rounding=False,
            ),
        )

        self.assertTrue(result["eligible"])
        self.assertTrue(result["margin_gap_review"])
        self.assertGreater(result["lifetime_margin_gap"], 0.05)
        self.assertIn(
            "Lifetime margin exceeds promo margin by more than 5 points",
            result["warnings"],
        )

    def test_lifetime_margin_gap_of_exactly_five_points_is_not_reviewed(self):
        row = {
            "sku": "CD1024-DGY-66X72-UK",
            "grade": 4,
            "cogs": 9.69,
            "avg_freight": 1.71,
            "ca_price": 25.99,
            "return_rate": 0,
            "lifetime_profit_margin": 0.17,
        }
        result = calculate_candidate(
            row,
            dict(
                CRITERIA,
                default_commission=0.24,
                price_source="ca",
                rounding=False,
            ),
        )

        self.assertFalse(result["margin_gap_review"])

    def test_lifetime_metrics_match_dashboard_all_platform_definitions(self):
        metrics = aggregate_lifetime_metrics(
            [
                {
                    "sku_code": "ABC-UK",
                    "sku_qty": 2,
                    "sales_amt": 100,
                    "extra_freight": 5,
                    "promo_rebate": 10,
                    "refund_amt": 4,
                    "resend_amt": 1,
                    "profit_incl_rn": 19,
                },
                {
                    "sku_code": "abc-uk",
                    "sku_qty": 1,
                    "sales_amt": 50,
                    "extra_freight": 0,
                    "promo_rebate": 0,
                    "refund_amt": 0,
                    "resend_amt": 5,
                    "profit_incl_rn": 5,
                },
            ]
        )["ABC-UK"]

        self.assertEqual(metrics["sold_qty"], 3)
        self.assertEqual(metrics["net_sales"], 145)
        self.assertAlmostEqual(metrics["return_rate"], 10 / 145)
        self.assertAlmostEqual(metrics["lifetime_profit_margin"], 24 / 145)
        self.assertEqual(
            metrics["normal_margin"],
            metrics["lifetime_profit_margin"],
        )

    def test_suggested_freight_matches_dashboard_precedence(self):
        self.assertEqual(select_suggested_freight(4.5, 20, 3.5, 2.5), 4.5)
        self.assertEqual(select_suggested_freight(None, 6, 3.5, 2.5), 3.5)
        self.assertEqual(select_suggested_freight(None, 5, 3.5, 2.5), 2.5)

    def test_reverse_pricing_includes_wms_and_hits_target_before_rounding(self):
        row = {
            "sku": "CD1024-DGY-66X72-UK",
            "grade": 4,
            "cogs": 9.69,
            "avg_freight": 1.71,
            "ca_price": 25.99,
        }
        result = calculate_candidate(
            row,
            dict(
                CRITERIA,
                default_commission=0.24,
                price_source="ca",
                rounding=False,
            ),
        )

        self.assertAlmostEqual(result["suggested_discount"], 0.06007805199802141)
        self.assertAlmostEqual(result["promo_margin"], 0.12)
        self.assertAlmostEqual(
            result["wms_fee"],
            result["promo_price_excluding_vat"] * 0.08,
        )

    def test_normal_margin_gap_does_not_create_review_warning(self):
        row = {
            "sku": "CD1024-DGY-66X72-UK",
            "grade": 4,
            "cogs": 9.69,
            "avg_freight": 1.71,
            "ca_price": 25.99,
            "sold_qty": 100,
            "normal_margin": 0.50,
        }
        result = calculate_candidate(
            row,
            dict(
                CRITERIA,
                default_commission=0.24,
                price_source="ca",
                rounding=False,
            ),
        )

        self.assertAlmostEqual(result["promo_margin"], result["target_margin"])
        self.assertFalse(
            any("normal margin" in warning.lower() for warning in result["warnings"])
        )

    def test_discount_interval_rounds_suggestion_down(self):
        row = {
            "sku": "CD1024-DGY-66X72-UK",
            "grade": 4,
            "cogs": 9.69,
            "avg_freight": 1.71,
            "ca_price": 25.99,
        }
        result = calculate_candidate(
            row,
            dict(
                CRITERIA,
                default_commission=0.24,
                price_source="ca",
                rounding=False,
                use_discount_interval=True,
                discount_interval=0.05,
            ),
        )

        self.assertAlmostEqual(result["suggested_discount"], 0.05)
        self.assertAlmostEqual(result["final_discount"], 0.05)
        self.assertGreater(result["promo_margin"], result["target_margin"])

    def test_discount_interval_does_not_change_manual_override(self):
        row = dict(
            self.rows[1],
            override_discount=0.071,
        )
        result = calculate_candidate(
            row,
            dict(
                CRITERIA,
                rounding=False,
                use_discount_interval=True,
                discount_interval=0.05,
            ),
        )

        self.assertAlmostEqual(result["final_discount"], 0.071)

    def test_commission_workbook_reads_all_platform_tabs(self):
        workbook = Workbook()
        range_sheet = workbook.active
        range_sheet.title = "The Range"
        range_sheet.append(["Platform", "Category", "SKU", "Commission Rate"])
        range_sheet.append(["The Range", "Dehumidifiers", "AI1005-BK-UK", 0.17])
        debenhams_sheet = workbook.create_sheet("Debenhams")
        debenhams_sheet.append(
            ["Platform", "Category", "SKU", "Commission Rate"]
        )
        debenhams_sheet.append(
            ["Debenhams", "Electricals", "AI1005-BK-UK", "13.2%"]
        )
        duplicate = ["Debenhams", "Electricals", "AI1005-BK-UK", 0.132]
        debenhams_sheet.append(duplicate)
        buffer = io.BytesIO()
        workbook.save(buffer)

        rows, conflicts = commission_rows_from_workbook(buffer.getvalue())

        self.assertFalse(conflicts)
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            {
                (row["platform"], row["sku"]): row["commission"]
                for row in rows
            },
            {
                ("The Range", "AI1005-BK-UK"): 0.17,
                ("Debenhams", "AI1005-BK-UK"): 0.132,
            },
        )

    def test_commission_workbook_rejects_conflicting_duplicate_rates(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Debenhams"
        sheet.append(["Platform", "Category", "SKU", "Commission Rate"])
        sheet.append(["Debenhams", "Electricals", "AI1005-BK-UK", 0.132])
        sheet.append(["Debenhams", "Electricals", "AI1005-BK-UK", 0.24])
        buffer = io.BytesIO()
        workbook.save(buffer)

        rows, conflicts = commission_rows_from_workbook(buffer.getvalue())

        self.assertEqual(len(rows), 1)
        self.assertEqual(len(conflicts), 1)
        self.assertEqual(conflicts[0]["sku"], "AI1005-BK-UK")

    def test_unconfirmed_commission_is_flagged_for_review(self):
        result = calculate_candidate(
            dict(self.rows[1], commission_requires_review=True),
            CRITERIA,
        )

        self.assertIn("Commission rate needs confirmation", result["warnings"])

    def test_ca_price_column_c_mapping_rule(self):
        self.assertEqual(ca_price_sku("AI1005-BK"), "AI1005-BK-UK")
        self.assertEqual(ca_price_sku("ABC-UK_offer-1"), "ABC-UK")
        self.assertEqual(ca_price_sku("ABC_BLUE"), "ABC-UK")
        self.assertEqual(ca_price_sku("AB.12"), "ABD12-UK")

    def test_mapping_resolver_only_prompts_after_all_automatic_methods_fail(self):
        inventory = {"EXACT", "RULE-UK", "SAVED-UK", "EXPLICIT-UK"}
        mappings = {"ODD-PLATFORM-SKU": "SAVED-UK"}

        self.assertEqual(resolve_wooper_sku("EXACT", inventory), "EXACT")
        self.assertEqual(resolve_wooper_sku("RULE", inventory), "RULE-UK")
        self.assertEqual(
            resolve_wooper_sku("ODD-PLATFORM-SKU", inventory, mappings),
            "SAVED-UK",
        )
        self.assertEqual(
            resolve_wooper_sku(
                "ANOTHER-PLATFORM-SKU",
                inventory,
                explicit_wooper_sku="EXPLICIT-UK",
            ),
            "EXPLICIT-UK",
        )
        self.assertIsNone(resolve_wooper_sku("UNKNOWN", inventory, mappings))

    def test_channeladvisor_prices_align_to_verified_wooper_skus(self):
        records, aligned_prices = align_channeladvisor_prices(
            {"EXACT", "RULE-UK", "SAVED-UK"},
            mappings={
                "ODD-PLATFORM-SKU": "SAVED-UK",
                "RETIRED-SKU": NON_EXISTING_SKU,
                "PARENT-ALL": "SAVED-UK",
            },
            exact_prices={
                "EXACT": 10,
                "RULE": 20,
                "ODD-PLATFORM-SKU": 30,
                "RETIRED-SKU": 40,
                "UNKNOWN": 50,
                "PARENT-ALL": 60,
            },
        )

        self.assertEqual(records["EXACT"]["mapping_source"], "exact")
        self.assertEqual(records["RULE"]["wooper_sku"], "RULE-UK")
        self.assertEqual(records["RULE"]["mapping_source"], "rule")
        self.assertEqual(records["ODD-PLATFORM-SKU"]["wooper_sku"], "SAVED-UK")
        self.assertEqual(records["ODD-PLATFORM-SKU"]["mapping_source"], "saved")
        self.assertEqual(
            records["RETIRED-SKU"]["mapping_status"],
            "non_existing",
        )
        self.assertEqual(records["UNKNOWN"]["mapping_status"], "unresolved")
        self.assertEqual(records["PARENT-ALL"]["mapping_status"], "non_existing")
        self.assertIsNone(records["PARENT-ALL"]["wooper_sku"])
        self.assertEqual(
            records["PARENT-ALL"]["mapping_source"],
            "parent_sku_rule",
        )
        self.assertEqual(
            aligned_prices,
            {"EXACT": 10, "RULE-UK": 20, "SAVED-UK": 30},
        )

    def test_imported_alias_can_use_ca_price_aligned_through_wooper_sku(self):
        mapped = map_imported_rows(
            [{"Platform SKU": "PROMO-ALIAS"}],
            manual_mappings={"PROMO-ALIAS": "CA-RULE-UK"},
            inventory={
                "CA-RULE-UK": {
                    "sku": "CA-RULE-UK",
                    "grade": 3,
                    "stock": 5,
                    "cogs": 8,
                }
            },
            exact_ca_prices={"CA-RULE": 42},
        )

        self.assertEqual(mapped[0]["sku"], "CA-RULE-UK")
        self.assertEqual(mapped[0]["ca_price"], 42)

    def test_non_existing_sku_is_removed_before_calculation(self):
        rows = [
            {"Platform SKU": "RETIRED-SKU", "Current Offer Price": 10},
            {"Platform SKU": "ACTIVE-UK", "Current Offer Price": 20},
        ]
        inventory = {
            "ACTIVE-UK": {
                "sku": "ACTIVE-UK",
                "grade": 3,
                "stock": 5,
                "cogs": 8,
            }
        }
        mapped = map_imported_rows(
            rows,
            manual_mappings={"RETIRED-SKU": NON_EXISTING_SKU},
            inventory=inventory,
        )

        self.assertEqual(len(mapped), 1)
        self.assertEqual(mapped[0]["sku"], "ACTIVE-UK")

    def test_inventory_sample_price_does_not_fill_a_missing_offer_price(self):
        mapped = map_imported_rows(
            [{"Platform SKU": "ACTIVE-UK"}],
            manual_mappings={},
            inventory={
                "ACTIVE-UK": {
                    "sku": "ACTIVE-UK",
                    "price": 99,
                    "grade": 3,
                    "stock": 5,
                    "cogs": 8,
                }
            },
        )

        self.assertEqual(len(mapped), 1)
        self.assertIsNone(mapped[0].get("offer_price"))
        self.assertNotIn("price", mapped[0])

    def test_first_arrival_cutoff_excludes_new_products(self):
        row = dict(self.rows[1], first_arrival_date="2026-07-01")
        criteria = dict(CRITERIA, exclude_first_arrival_on_or_after="2026-06-01")
        actual = calculate_candidate(row, criteria)
        self.assertFalse(actual["eligible"])
        self.assertIn(
            "First arrival is inside the excluded new-product period",
            actual["reasons"],
        )

    def test_first_arrival_cutoff_normalises_uk_dates(self):
        row = dict(self.rows[1], first_arrival_date="23/7/2026")
        criteria = dict(CRITERIA, exclude_first_arrival_on_or_after="1/6/2026")
        actual = calculate_candidate(row, criteria)

        self.assertEqual(normalise_date("23/7/2026"), "2026-07-23")
        self.assertEqual(normalise_date("1/6/2026"), "2026-06-01")
        self.assertFalse(actual["eligible"])
        self.assertIn(
            "First arrival is inside the excluded new-product period",
            actual["reasons"],
        )


if __name__ == "__main__":
    unittest.main()
