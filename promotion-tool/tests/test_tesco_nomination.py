import io
import json
import sys
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app


class TescoNominationTests(unittest.TestCase):
    def test_candidate_import_only_accepts_tesco_approval_yes(self):
        rows = [
            {"Platform": "Tesco", "Platform SKU": "ABC_1", "SKU": "ABC-UK", "Approval": "Yes", "SOH": 7},
            {"Platform": "Tesco", "Platform SKU": "ABC_1", "SKU": "ABC-UK", "Approval": "Yes", "SOH": 7},
            {"Platform": "Tesco", "Platform SKU": "DEF_1", "SKU": "DEF-UK", "Approval": "No"},
            {"Platform": "eBay", "Platform SKU": "GHI", "SKU": "GHI-UK", "Approval": "Yes"},
        ]
        candidates = app.tesco_candidates_from_rows(rows, inventory={"ABC-UK": {"grade": 4, "estimated_months": 2.5}})
        self.assertEqual([row["tesco_sku"] for row in candidates], ["ABC_1"])
        self.assertEqual(candidates[0]["soh"], 7)
        self.assertEqual(candidates[0]["grade"], 4)
        self.assertEqual(candidates[0]["months"], 2.5)

    def test_overlap_is_by_tesco_sku_not_wooper_sku(self):
        events = [{
            "event_name": "Event A", "start_date": "2026-08-01", "end_date": "2026-08-20",
            "rows": [{"tesco_sku": "HO0603-3-PK-UK", "wooper_sku": "HO0603-3-PK-UK"}],
        }]
        with patch.object(app, "tesco_events", return_value=events):
            conflicts = app.tesco_conflicts("2026-08-18", "2026-09-01")
        self.assertIn("HO0603-3-PK-UK", conflicts)
        self.assertNotIn("HO0603-3-PK-UK_1", conflicts)

    def test_catalogue_sku_uses_platform_to_wooper_mapping(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Category Code", "SKU", "Title", "Brand", "Barcode", "Image 1"])
        sheet.append(["Home/Furniture", "FR2044-BR", "Brown rug", "Levede", "123", "https://example.com/image.jpg"])
        output = io.BytesIO()
        workbook.save(output)
        rows = app.tesco_catalogue_from_workbook(
            output.getvalue(),
            inventory={"FR2044-BR-UK": {}},
        )
        self.assertEqual(rows[0]["tesco_sku"], "FR2044-BR")
        self.assertEqual(rows[0]["wooper_sku"], "FR2044-BR-UK")

    def test_external_nomination_matches_unique_catalogue_barcode(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Seller Input"])
        sheet.append(["Event line", "Start Date", "End Date", "Barcode/EAN", "Product Title"])
        sheet.append(["Input if for a specific event", "01/01/2026", "02/01/2026", "1234567890", "Panasonic 50-inch TV"])
        sheet.append(["Autumn Event", "01/09/2026", "14/09/2026", 1234567890123, "Example item"])
        output = io.BytesIO()
        workbook.save(output)
        result = app.tesco_external_nomination_from_workbook(
            output.getvalue(),
            catalogue=[{
                "tesco_sku": "TESCO-123", "wooper_sku": "CORE-123-UK",
                "barcode": "1234567890123", "title": "Example item",
            }],
        )
        event = result["events"][0]
        self.assertEqual(event["event_name"], "Autumn Event")
        self.assertEqual(event["start_date"], "2026-09-01")
        self.assertEqual(len(event["rows"]), 1)
        self.assertEqual(event["rows"][0]["tesco_sku"], "TESCO-123")
        self.assertEqual(event["rows"][0]["match_source"], "barcode")

    def test_external_nomination_leaves_ambiguous_barcode_for_confirmation(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Event", "Start Date", "End Date", "Barcode / EAN", "Product Title"])
        sheet.append(["Input if for a specific event", "2026-01-01", "2026-01-02", "1234567890", "Panasonic 50-inch TV"])
        sheet.append(["Duplicate Listing Event", "2026-09-01", "2026-09-14", "123", "Example item"])
        output = io.BytesIO()
        workbook.save(output)
        result = app.tesco_external_nomination_from_workbook(
            output.getvalue(),
            catalogue=[
                {"tesco_sku": "TESCO-123", "wooper_sku": "CORE-UK", "barcode": "123"},
                {"tesco_sku": "TESCO-123_1", "wooper_sku": "CORE-UK", "barcode": "123"},
            ],
        )
        row = result["events"][0]["rows"][0]
        self.assertEqual(row["tesco_sku"], "")
        self.assertEqual(row["match_status"], "ambiguous")
        self.assertEqual([item["tesco_sku"] for item in row["suggestions"]], ["TESCO-123", "TESCO-123_1"])

    def test_internal_record_preserves_exact_tesco_listing(self):
        payload = {
            "event_name": "Saved Form", "start_date": "2026-09-01", "end_date": "2026-09-14",
            "rows": [{
                "tesco_sku": "TESCO-123_1", "wooper_sku": "CORE-UK", "brand": "Brand",
                "level_1": "Home", "event_category": "Home", "barcode": "123",
                "title": "Example", "ca_price": 20, "promo_price": 15, "soh": 1, "discount": .25,
            }],
        }
        with patch.object(app, "tesco_conflicts", return_value={}):
            workbook = app.build_tesco_nomination_workbook(payload)
        result = app.tesco_external_nomination_from_workbook(
            workbook,
            catalogue=[
                {"tesco_sku": "TESCO-123", "wooper_sku": "CORE-UK", "barcode": "123"},
                {"tesco_sku": "TESCO-123_1", "wooper_sku": "CORE-UK", "barcode": "123"},
            ],
        )
        row = result["events"][0]["rows"][0]
        self.assertEqual(row["tesco_sku"], "TESCO-123_1")
        self.assertEqual(row["match_source"], "internal_record")

    def test_external_event_save_records_confirmed_exact_tesco_sku(self):
        payload = {"events": [{
            "event_name": "External Event", "start_date": "2026-10-01", "end_date": "2026-10-14",
            "rows": [{"tesco_sku": "TESCO-123_1", "wooper_sku": "", "barcode": "123", "title": "Example"}],
        }]}
        with patch.object(app, "latest_tesco_catalogue", return_value=[{
            "tesco_sku": "TESCO-123_1", "wooper_sku": "CORE-UK", "barcode": "123",
        }]), patch.object(app, "tesco_events", return_value=[]), patch.object(
            app, "save_tesco_event", side_effect=lambda event, source=None: {**event, "source": source}
        ):
            saved = app.import_external_tesco_events(payload, "submitted.xlsx")
        self.assertEqual(saved[0]["rows"][0]["tesco_sku"], "TESCO-123_1")
        self.assertEqual(saved[0]["rows"][0]["wooper_sku"], "CORE-UK")
        self.assertTrue(saved[0]["source"]["imported_outside_tool"])

    def test_generated_form_uses_ca_price_stock_plus_100_and_required_constants(self):
        payload = {
            "event_name": "Summer Sale", "start_date": "2026-06-30", "end_date": "2026-09-01",
            "seller_name": "Traderight Group",
            "rows": [{
                "tesco_sku": "ABC_1", "wooper_sku": "ABC-UK", "brand": "Levede",
                "level_1": "Home", "event_category": "Home Comfort", "category_path": "Home > Bedding", "barcode": "1234567890123",
                "title": "Example product", "ca_price": 49.99, "promo_price": 39.99, "soh": 25,
                "discount": 0.20,
            }],
        }
        with patch.object(app, "tesco_conflicts", return_value={}):
            data = app.build_tesco_nomination_workbook(payload)
        workbook = load_workbook(io.BytesIO(data), data_only=False)
        sheet = workbook["Sheet1"]
        self.assertEqual(sheet["A1"].value, "Seller Input")
        self.assertEqual(sheet["A2"].value, "Seller Name")
        self.assertEqual(sheet["E3"].value, "Home > Bedding")
        self.assertEqual(sheet["I3"].value, None)
        self.assertEqual(sheet["J3"].value, None)
        self.assertEqual(sheet["L3"].value, 49.99)
        self.assertEqual(sheet["M3"].value, 39.99)
        self.assertEqual(sheet["N3"].value, "=L3-M3")
        self.assertEqual(sheet["P3"].value, "Exclusive Pricing")
        self.assertEqual(sheet["Q3"].value, 125)
        self.assertEqual([workbook["Sheet2"].cell(row, 1).value for row in range(1, 11)], app.TESCO_LEVEL_1)
        self.assertTrue(sheet.data_validations.dataValidation)
        internal = workbook["Internal Record"]
        self.assertEqual([internal.cell(1, column).value for column in range(1, 6)], [
            "Platform SKU",
            "Wooper SKU",
            "Original Price (CA Price, VAT Included)",
            "Discounted Price",
            "Discount %",
        ])
        self.assertEqual([internal.cell(2, column).value for column in range(1, 6)], [
            "ABC_1", "ABC-UK", 49.99, 39.99, 0.20,
        ])

    def test_generated_form_enforces_event_category_limit(self):
        payload = {
            "event_name": "Limited Event", "start_date": "2026-11-01", "end_date": "2026-11-10",
            "category_limits": [{"name": "Home", "limit": 1}],
            "rows": [
                {"tesco_sku": f"SKU-{index}", "event_category": "Home", "level_1": "Home"}
                for index in range(2)
            ],
        }
        with self.assertRaisesRegex(ValueError, "Home SKU limit"):
            app.build_tesco_nomination_workbook(payload)

    def test_saved_event_can_be_removed(self):
        root = Path(__file__).parent / ".tesco_event_delete_test"
        nominations = root / "nominations"
        events_path = root / "events.json"
        root.mkdir(exist_ok=True)
        nominations.mkdir(exist_ok=True)
        try:
            events_path.write_text(json.dumps({"events": [{"id": "event-1", "event_name": "Wrong event"}]}))
            (nominations / "event-1.xlsx").write_bytes(b"test")
            with patch.object(app, "TESCO_EVENTS_PATH", events_path), patch.object(app, "TESCO_NOMINATIONS_DIR", nominations):
                removed = app.delete_tesco_event("event-1")
            self.assertEqual(removed["event_name"], "Wrong event")
            self.assertEqual(json.loads(events_path.read_text())["events"], [])
            self.assertFalse((nominations / "event-1.xlsx").exists())
        finally:
            if events_path.exists():
                events_path.unlink()
            if nominations.exists():
                nominations.rmdir()
            if root.exists():
                root.rmdir()

    def test_saved_event_mirakl_text_can_be_updated_and_cleared(self):
        root = Path(__file__).parent / ".tesco_event_update_test"
        events_path = root / "events.json"
        root.mkdir(exist_ok=True)
        try:
            events_path.write_text(json.dumps({"events": [{"id": "event-1", "event_name": "Event"}]}))
            with patch.object(app, "TESCO_EVENTS_PATH", events_path), patch.object(app, "supabase_enabled", return_value=False):
                updated = app.update_tesco_event("event-1", {"discount_end_date_mirakl": "30 Sep - extended"})
                saved_value = updated["discount_end_date_mirakl"]
                cleared = app.update_tesco_event("event-1", {"discount_end_date_mirakl": ""})
            self.assertEqual(saved_value, "30 Sep - extended")
            self.assertIsNone(cleared["discount_end_date_mirakl"])
            self.assertIsNone(json.loads(events_path.read_text())["events"][0]["discount_end_date_mirakl"])
        finally:
            if events_path.exists():
                events_path.unlink()
            if root.exists():
                root.rmdir()

    def test_tesco_offer_capture_uses_shared_supabase_storage(self):
        source_rows = [{"Platform SKU": "ABC_1", "Buy It Now Price": 19.99}]
        mapped_rows = [{
            "platform_sku": "ABC_1", "sku": "ABC-UK", "ca_price": 19.99,
            "mapping_status": "mapped",
        }]
        with patch.object(app, "supabase_enabled", return_value=True), patch.object(
            app, "supabase_request", return_value=[]
        ) as request:
            offers = app.capture_tesco_offers(source_rows, mapped_rows)
        self.assertEqual(offers[0]["wooper_sku"], "ABC-UK")
        self.assertEqual(request.call_args_list[0].args[:2], ("POST", "promotion_tesco_offers"))
        self.assertEqual(request.call_args_list[1].args[:2], ("DELETE", "promotion_tesco_offers"))

    def test_empty_supabase_tesco_history_is_seeded_from_local_records(self):
        root = Path(__file__).parent / ".tesco_event_seed_test"
        events_path = root / "events.json"
        root.mkdir(exist_ok=True)
        try:
            events_path.write_text(json.dumps({"events": [{
                "id": "promotion-master-a", "event_name": "Existing event",
                "start_date": "2026-01-01", "end_date": "2026-01-10",
                "source": "Promotion Master.xlsx", "rows": [{"tesco_sku": "ABC"}],
            }]}))
            with patch.object(app, "TESCO_EVENTS_PATH", events_path), patch.object(
                app, "supabase_enabled", return_value=True
            ), patch.object(app, "supabase_select_all", return_value=[]), patch.object(
                app, "supabase_request", side_effect=lambda *args, **kwargs: kwargs["rows"]
            ) as request:
                events = app.tesco_events()
            self.assertEqual(events[0]["event_name"], "Existing event")
            self.assertEqual(events[0]["source"]["file"], "Promotion Master.xlsx")
            self.assertEqual(request.call_args.args[:2], ("POST", "promotion_tesco_events"))
            self.assertEqual(str(uuid.UUID(events[0]["id"])), events[0]["id"])
        finally:
            if events_path.exists():
                events_path.unlink()
            if root.exists():
                root.rmdir()


if __name__ == "__main__":
    unittest.main()
