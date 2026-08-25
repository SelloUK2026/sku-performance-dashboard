from __future__ import annotations

import csv
import base64
import hmac
import io
import json
import math
import os
import re
import time
import uuid
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Lock
from urllib.error import HTTPError
from urllib.parse import parse_qs, urlencode, urlparse
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
SAMPLE_PATH = BASE_DIR / "sample_data.json"
DATA_DIR = BASE_DIR / "data"
MAPPINGS_PATH = BASE_DIR / "data" / "sku_mappings.json"
LIFETIME_METRICS_PATH = BASE_DIR / "data" / "lifetime_metrics.json"
TESCO_OFFERS_PATH = DATA_DIR / "tesco_latest_offers.json"
TESCO_CATALOGUE_PATH = DATA_DIR / "tesco_latest_catalogue.json"
TESCO_EVENTS_PATH = DATA_DIR / "tesco_nomination_events.json"
TESCO_NOMINATIONS_DIR = DATA_DIR / "tesco_nominations"
PLATFORM_SETTINGS_PATH = DATA_DIR / "platform_settings.json"
NON_EXISTING_SKU = "__NON_EXISTING__"
WORKBOOK_PATH = Path(
    os.environ.get(
        "SKU_DATA_WORKBOOK",
        BASE_DIR.parent / "Lastest Data Analyse - Codex.xlsx",
    )
)
WOOPER_INVENTORY_PATH = os.environ.get("WOOPER_INVENTORY_PATH")
CHANNELADVISOR_PATH = os.environ.get("CHANNELADVISOR_PATH")
SUPABASE_URL = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
if SUPABASE_URL.endswith("/rest/v1"):
    SUPABASE_URL = SUPABASE_URL[:-8].rstrip("/")
SUPABASE_SERVICE_ROLE_KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
PROMOTION_DATA_SOURCE = os.environ.get(
    "PROMOTION_DATA_SOURCE",
    "supabase" if SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY else "workbook",
).strip().lower()
PROMOTION_CACHE_SECONDS = int(os.environ.get("PROMOTION_CACHE_SECONDS", "300"))
PROMOTION_AUTH_USERNAME = os.environ.get("PROMOTION_AUTH_USERNAME", "").strip()
PROMOTION_AUTH_PASSWORD = os.environ.get("PROMOTION_AUTH_PASSWORD", "")
PROMOTION_REQUIRE_AUTH = (
    os.environ.get("PROMOTION_REQUIRE_AUTH", "false").strip().lower()
    in {"1", "true", "yes", "on"}
)
_INVENTORY_CACHE = {"path": None, "mtime_ns": None, "rows": None}
_FIRST_ARRIVAL_CACHE = {"path": None, "mtime_ns": None, "rows": None}
_FREIGHT_CACHE = {"path": None, "mtime_ns": None, "rows": None}
_PERFORMANCE_CACHE = {"path": None, "mtime_ns": None, "rows": None}
_PERFORMANCE_BUILD_LOCK = Lock()
_CA_PRICE_CACHE = {
    "path": None,
    "mtime_ns": None,
    "exact": None,
    "canonical": None,
}
_SUPABASE_PROMOTION_CACHE = {"expires_at": 0.0, "rows": None}
_SUPABASE_CA_PRICE_CACHE = {
    "expires_at": 0.0,
    "exact": None,
    "canonical": None,
}
WORKTABLE_SCHEMA_VERSION = 1
MAX_WORKTABLE_SNAPSHOT_BYTES = 12 * 1024 * 1024
TESCO_LEVEL_1 = [
    "Baby",
    "BWS",
    "Celebration",
    "Electricals",
    "Health & Beauty",
    "Home",
    "Outdoor",
    "Pet",
    "Sports & Leisure",
    "Toys",
]

DEFAULT_COMMISSIONS = {
    "eBay": 0.11,
    "Amazon(UK)": 0.18,
    "Temu(UK)": 0.00,
    "Wayfair": 0.05,
    "Debenhams": 0.24,
    "Tesco": 0.18,
    "BrandAlley": 0.24,
    "Decathlon UK Limited": 0.19,
    "The Range": 0.14,
    "TikTok(Skylos)": 0.09,
    "Tiktok(Levede)": 0.09,
    "Skylous shopify": 0.00,
    "Go Groopie": 0.00,
    "Groupon(UK)": 0.00,
    "Wowcher": 0.20,
    "Onbuy": 0.15,
    "ManoMano": 0.16,
    "Fruugo": 0.20,
    "Rackham": 0.18,
}
VARIABLE_COMMISSION_PLATFORMS = {"Debenhams", "The Range"}


def default_platform_settings():
    return [
        {
            "platform": platform,
            "default_commission": commission,
            "manual_promo_price_adjustment": False,
        }
        for platform, commission in DEFAULT_COMMISSIONS.items()
    ]


ALIASES = {
    "platform_sku": {
        "platform sku",
        "platform_sku",
        "merchant sku",
        "seller sku",
        "sku",
        "inventory number",
    },
    "product_id": {"product-id", "product id", "product_id", "offer id"},
    "sku": {"sku code", "sku_code"},
    "wooper_sku": {"wooper id", "wooper sku", "core sku"},
    "brand": {"brand"},
    "subcategory": {"subcategory", "sub category", "category"},
    "inventory_status": {"inventory status", "inventory_status", "status"},
    "grade": {"grade", "grade level", "grade_level"},
    "estimated_months": {
        "estimated sales month",
        "estimated months to sell",
        "estimated_months",
        "saleable months",
    },
    "sold_qty": {"soldqty", "sold qty", "sold_qty", "quantity sold"},
    "sales_amt": {"salesamt", "sales amt", "sales_amt", "sales"},
    "postage": {"postage", "shipping cost"},
    "return_rate": {"res&ref%", "return rate", "return_rate", "resend refund rate"},
    "normal_margin": {"profit% before res+ref", "normal margin", "normal_margin"},
    "offer_price": {
        "price",
        "selling price",
        "offer price",
        "current offer price",
    },
    "ca_price": {
        "buy it now price",
        "buy now price",
        "bin price",
        "ca price",
        "normal price",
    },
    "cogs": {"cogs", "cost", "cost price"},
    "stock": {"soh", "stock", "stock on hand", "total inventory qty"},
    "commission": {"commission", "platform commission", "commission rate"},
    "avg_freight": {"avg freight", "average freight", "avg_freight", "suggested freight"},
    "override_discount": {"override discount", "override_discount"},
}


def clean_header(value) -> str:
    text = re.sub(r"[\r\n_]+", " ", str(value or "")).strip().lower()
    return re.sub(r"\s+", " ", text)


def normalise_number(value, default=0.0):
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return default if not math.isfinite(float(value)) else float(value)
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1]
        try:
            return float(text) / 100
        except ValueError:
            return default
    try:
        parsed = float(text)
        return parsed if math.isfinite(parsed) else default
    except ValueError:
        return default


def nullable_number(value):
    if value is None or value == "":
        return None
    return normalise_number(value, None)


def normalise_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    text = str(value).strip()
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text[:10], date_format).date().isoformat()
        except ValueError:
            continue
    return None


def excel_mround(value: float, multiple: float) -> float:
    if multiple == 0:
        return value
    quotient = Decimal(str(value)) / Decimal(str(multiple))
    rounded = quotient.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return float(rounded * Decimal(str(multiple)))


def grade_key(value) -> str:
    grade = int(normalise_number(value, 0))
    return str(grade if 0 <= grade <= 7 else 0)


def ca_price_sku(platform_sku) -> str:
    """Match the workbook formula in CA Price column C."""
    text = str(platform_sku or "").strip()
    upper = text.upper()
    uk_index = upper.find("-UK")
    if uk_index >= 0:
        mapped = text[: uk_index + 3]
    elif "_" in text:
        mapped = text.split("_", 1)[0] + "-UK"
    else:
        mapped = text + "-UK"
    return mapped.replace(".", "D").upper()


def load_mappings():
    if not MAPPINGS_PATH.exists():
        return {}
    return json.loads(MAPPINGS_PATH.read_text(encoding="utf-8"))


def save_mappings(mappings):
    MAPPINGS_PATH.parent.mkdir(parents=True, exist_ok=True)
    MAPPINGS_PATH.write_text(
        json.dumps(mappings, indent=2, ensure_ascii=True), encoding="utf-8"
    )


def supabase_enabled():
    return bool(SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY)


def supabase_request(method, table, rows=None, params=None, prefer=None):
    query = urlencode(params or {}, doseq=True)
    url = f"{SUPABASE_URL}/rest/v1/{table}"
    if query:
        url = f"{url}?{query}"
    headers = {
        "apikey": SUPABASE_SERVICE_ROLE_KEY,
        "Authorization": f"Bearer {SUPABASE_SERVICE_ROLE_KEY}",
        "Content-Type": "application/json",
    }
    if prefer:
        headers["Prefer"] = prefer
    body = None if rows is None else json.dumps(rows).encode("utf-8")
    request = Request(url, data=body, headers=headers, method=method)
    try:
        with urlopen(request, timeout=120) as response:
            payload = response.read()
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(
            f"Supabase {method} {table} failed: {exc.code} {detail}"
        ) from exc
    return json.loads(payload.decode("utf-8")) if payload else []


def supabase_select_all(table, params=None, page_size=1000):
    rows = []
    offset = 0
    while True:
        page_params = dict(params or {})
        page_params["limit"] = page_size
        page_params["offset"] = offset
        page = supabase_request("GET", table, params=page_params)
        rows.extend(page)
        if len(page) < page_size:
            return rows
        offset += page_size


def validate_platform_setting(row):
    if not isinstance(row, dict):
        raise ValueError("Every platform setting must be an object.")
    platform = str(row.get("platform") or "").strip()
    if not platform:
        raise ValueError("Platform name is required.")
    if len(platform) > 120:
        raise ValueError("Platform name is too long.")
    commission = normalise_number(row.get("default_commission"), None)
    if commission is None or commission < 0 or commission > 1:
        raise ValueError("Default commission must be between 0% and 100%.")
    return {
        "platform": platform,
        "default_commission": commission,
        "manual_promo_price_adjustment": bool(
            row.get("manual_promo_price_adjustment", False)
        ),
    }


def platform_settings():
    if supabase_enabled():
        rows = supabase_request(
            "GET",
            "promotion_platform_settings",
            params={
                "select": (
                    "platform,default_commission,"
                    "manual_promo_price_adjustment"
                ),
                "order": "platform.asc",
            },
        )
        return [validate_platform_setting(row) for row in rows]
    if PLATFORM_SETTINGS_PATH.exists():
        rows = json.loads(PLATFORM_SETTINGS_PATH.read_text(encoding="utf-8"))
        return [validate_platform_setting(row) for row in rows]
    return default_platform_settings()


def save_platform_settings(payload):
    rows = payload.get("settings") if isinstance(payload, dict) else None
    if not isinstance(rows, list) or not rows:
        raise ValueError("Add at least one platform setting.")
    settings = [validate_platform_setting(row) for row in rows]
    names = [row["platform"].casefold() for row in settings]
    if len(names) != len(set(names)):
        raise ValueError("Platform names must be unique.")
    if supabase_enabled():
        supabase_request(
            "POST",
            "promotion_platform_settings",
            rows=settings,
            params={"on_conflict": "platform"},
            prefer="resolution=merge-duplicates,return=representation",
        )
    else:
        PLATFORM_SETTINGS_PATH.parent.mkdir(parents=True, exist_ok=True)
        PLATFORM_SETTINGS_PATH.write_text(
            json.dumps(settings, indent=2, ensure_ascii=True),
            encoding="utf-8",
        )
    return platform_settings()


def validate_worktable_id(value):
    try:
        return str(uuid.UUID(str(value or "")))
    except (ValueError, TypeError, AttributeError) as exc:
        raise ValueError("Invalid saved worktable ID.") from exc


def list_saved_worktables(platform="", event_name="", created_on="", limit=100):
    if not supabase_enabled():
        raise RuntimeError("Saved worktables require Supabase.")
    params = {
        "select": (
            "id,platform,event_name,source_file,source_row_count,"
            "candidate_count,eligible_count,selected_count,created_at,created_on"
        ),
        "order": "created_at.desc",
        "limit": max(1, min(int(limit or 100), 250)),
    }
    platform = str(platform or "").strip()
    event_name = str(event_name or "").strip()
    created_on = str(created_on or "").strip()
    if platform:
        params["platform"] = f"eq.{platform}"
    if event_name:
        params["event_name"] = f"ilike.*{event_name[:160]}*"
    if created_on:
        try:
            date.fromisoformat(created_on)
        except ValueError as exc:
            raise ValueError("Creation date must use YYYY-MM-DD.") from exc
        params["created_on"] = f"eq.{created_on}"
    return supabase_request("GET", "promotion_worktables", params=params)


def get_saved_worktable(worktable_id):
    if not supabase_enabled():
        raise RuntimeError("Saved worktables require Supabase.")
    rows = supabase_request(
        "GET",
        "promotion_worktables",
        params={
            "select": "*",
            "id": f"eq.{validate_worktable_id(worktable_id)}",
            "limit": 1,
        },
    )
    return rows[0] if rows else None


def create_saved_worktable(payload):
    if not supabase_enabled():
        raise RuntimeError("Saved worktables require Supabase.")
    platform = str(payload.get("platform") or "").strip()
    event_name = str(payload.get("event_name") or "").strip()
    snapshot = payload.get("snapshot")
    if not platform:
        raise ValueError("Promotion platform is required.")
    if len(platform) > 120:
        raise ValueError("Promotion platform is too long.")
    if not event_name:
        raise ValueError("Event name is required.")
    if len(event_name) > 160:
        raise ValueError("Event name is too long.")
    if not isinstance(snapshot, dict):
        raise ValueError("Worktable snapshot is required.")
    candidates = snapshot.get("candidates")
    rows = snapshot.get("rows")
    selected_skus = snapshot.get("selected_skus")
    if not isinstance(candidates, list) or not candidates:
        raise ValueError("Calculate at least one SKU before saving.")
    if not isinstance(rows, list) or not isinstance(selected_skus, list):
        raise ValueError("Worktable snapshot is incomplete.")
    snapshot = dict(snapshot)
    snapshot["schema_version"] = WORKTABLE_SCHEMA_VERSION
    encoded_snapshot = json.dumps(
        snapshot,
        ensure_ascii=False,
        allow_nan=False,
        separators=(",", ":"),
    ).encode("utf-8")
    if len(encoded_snapshot) > MAX_WORKTABLE_SNAPSHOT_BYTES:
        raise ValueError("This worktable is too large to save.")
    source = snapshot.get("source") if isinstance(snapshot.get("source"), dict) else {}
    record = {
        "platform": platform,
        "event_name": event_name,
        "source_file": str(source.get("file") or "").strip() or None,
        "source_row_count": max(0, int(source.get("row_count") or len(rows))),
        "candidate_count": len(candidates),
        "eligible_count": sum(bool(row.get("eligible")) for row in candidates),
        "selected_count": len({str(sku) for sku in selected_skus if str(sku)}),
        "snapshot": snapshot,
    }
    saved = supabase_request(
        "POST",
        "promotion_worktables",
        rows=[record],
        prefer="return=representation",
    )
    if not saved:
        raise RuntimeError("Supabase did not return the saved worktable.")
    return saved[0]


def persisted_mapping_rows(
    mapping_scope,
    platform="",
    include_all_platforms=False,
):
    if not supabase_enabled():
        return []
    rows = supabase_select_all(
        "sku_mappings",
        {
            "select": "mapping_scope,platform,external_sku,wooper_sku,status",
            "mapping_scope": f"eq.{mapping_scope}",
        },
    )
    if mapping_scope == "channeladvisor":
        return [row for row in rows if not row.get("platform")]
    if include_all_platforms:
        return rows
    platform_key = str(platform or "").strip()
    return [
        row
        for row in rows
        if not row.get("platform") or row.get("platform") == platform_key
    ]


def persisted_mapping_dictionary(mapping_scope, platform=""):
    if not supabase_enabled():
        return load_mappings()
    mappings = {}
    rows = persisted_mapping_rows(mapping_scope, platform)
    rows.sort(key=lambda row: bool(row.get("platform")))
    for row in rows:
        external_sku = str(row.get("external_sku") or "").strip().upper()
        if not external_sku:
            continue
        mapped_value = (
            NON_EXISTING_SKU
            if row.get("status") == "non_existing"
            else str(row.get("wooper_sku") or "").strip().upper()
        )
        if mapped_value:
            mappings[external_sku] = mapped_value
    return mappings


def combined_platform_mappings(platform=""):
    if not supabase_enabled():
        return load_mappings()
    mappings = persisted_mapping_dictionary("channeladvisor")
    rows = persisted_mapping_rows(
        "platform",
        include_all_platforms=True,
    )
    grouped = {}
    for row in rows:
        external_sku = str(row.get("external_sku") or "").strip().upper()
        if not external_sku:
            continue
        mapped_value = (
            NON_EXISTING_SKU
            if row.get("status") == "non_existing"
            else str(row.get("wooper_sku") or "").strip().upper()
        )
        if mapped_value:
            grouped.setdefault(external_sku, set()).add(mapped_value)

    # Reuse an external SKU across marketplaces only when every saved row agrees.
    for external_sku, mapped_values in grouped.items():
        if len(mapped_values) == 1:
            mappings[external_sku] = next(iter(mapped_values))

    platform_key = str(platform or "").strip()
    current_rows = [
        row
        for row in rows
        if not row.get("platform") or row.get("platform") == platform_key
    ]
    current_rows.sort(key=lambda row: bool(row.get("platform")))
    for row in current_rows:
        external_sku = str(row.get("external_sku") or "").strip().upper()
        if not external_sku:
            continue
        mapped_value = (
            NON_EXISTING_SKU
            if row.get("status") == "non_existing"
            else str(row.get("wooper_sku") or "").strip().upper()
        )
        if mapped_value:
            mappings[external_sku] = mapped_value
    return mappings


def unresolved_channeladvisor_rows():
    if not supabase_enabled():
        return []
    return supabase_select_all(
        "channeladvisor_products",
        {
            "select": "platform_sku,title,brand,ca_price,mapping_status",
            "mapping_status": "eq.unresolved",
            "order": "platform_sku.asc",
        },
    )


def save_persisted_mappings(items, mapping_scope, platform=""):
    if not supabase_enabled():
        saved = load_mappings()
        for item in items:
            external_sku = item["external_sku"]
            saved[external_sku] = (
                NON_EXISTING_SKU
                if item["status"] == "non_existing"
                else item["wooper_sku"]
            )
        save_mappings(saved)
        return

    timestamp = datetime.now().astimezone().isoformat()
    rows = [
        {
            "mapping_scope": mapping_scope,
            "platform": "" if mapping_scope == "channeladvisor" else platform,
            "external_sku": item["external_sku"],
            "wooper_sku": item["wooper_sku"],
            "status": item["status"],
            "mapping_source": "manual",
            "updated_at": timestamp,
        }
        for item in items
    ]
    if not rows:
        return
    supabase_request(
        "POST",
        "sku_mappings",
        rows=rows,
        params={"on_conflict": "mapping_scope,platform,external_sku"},
        prefer="resolution=merge-duplicates,return=minimal",
    )
    if mapping_scope == "channeladvisor":
        product_rows = [
            {
                "platform_sku": item["external_sku"],
                "wooper_sku": item["wooper_sku"],
                "mapping_status": item["status"],
                "mapping_source": "manual",
            }
            for item in items
        ]
        supabase_request(
            "POST",
            "channeladvisor_products",
            rows=product_rows,
            params={"on_conflict": "platform_sku"},
            prefer="resolution=merge-duplicates,return=minimal",
        )


def supabase_inventory_skus():
    if not supabase_enabled():
        return []
    if PROMOTION_DATA_SOURCE == "supabase":
        return sorted(supabase_promotion_inventory())
    return sorted(
        str(row.get("sku") or "").strip().upper()
        for row in supabase_select_all("inventory", {"select": "sku"})
        if row.get("sku")
    )


def promotion_data_refreshed_at():
    if not supabase_enabled():
        return None
    rows = supabase_request(
        "GET",
        "promotion_sku_data",
        params={
            "select": "refreshed_at",
            "order": "refreshed_at.desc",
            "limit": 1,
        },
    )
    return rows[0].get("refreshed_at") if rows else None


def active_protection_rows(as_of=None):
    if not supabase_enabled():
        return []
    as_of = as_of or datetime.now(ZoneInfo("Australia/Sydney")).date()
    as_of_text = as_of.isoformat() if isinstance(as_of, date) else str(as_of)
    return supabase_select_all(
        "promotion_protection_list",
        {
            "select": (
                "sku,protection_owner,protected_ca_price,protection_start,"
                "protection_end,refreshed_at"
            ),
            "protection_start": f"lte.{as_of_text}",
            "protection_end": f"gte.{as_of_text}",
            "order": "sku.asc",
        },
    )


def supabase_promotion_inventory():
    if not supabase_enabled():
        return {}
    now = time.monotonic()
    if (
        _SUPABASE_PROMOTION_CACHE["rows"] is not None
        and now < _SUPABASE_PROMOTION_CACHE["expires_at"]
    ):
        return _SUPABASE_PROMOTION_CACHE["rows"]

    rows = supabase_select_all(
        "promotion_sku_data",
        {
            "select": (
                "sku,main_category,subcategory,brand,inventory_status,"
                "grade_level,estimated_months_to_sell,stock_on_hand,cogs,"
                "first_arrival_date,suggested_freight,sold_qty,sales_amt,"
                "return_rate,lifetime_profit_margin"
            ),
            "order": "sku.asc",
        },
    )
    inventory = {}
    for row in rows:
        sku = str(row.get("sku") or "").strip().upper()
        if not sku:
            continue
        inventory[sku] = {
            "sku": sku,
            "main_category": row.get("main_category"),
            "subcategory": row.get("subcategory"),
            "brand": row.get("brand"),
            "inventory_status": row.get("inventory_status"),
            "grade": nullable_number(row.get("grade_level")),
            "estimated_months": nullable_number(
                row.get("estimated_months_to_sell")
            ),
            "stock": nullable_number(row.get("stock_on_hand")),
            "cogs": nullable_number(row.get("cogs")),
            "first_arrival_date": row.get("first_arrival_date"),
            "avg_freight": nullable_number(row.get("suggested_freight")),
            "sold_qty": normalise_number(row.get("sold_qty")),
            "sales_amt": normalise_number(row.get("sales_amt")),
            "return_rate": nullable_number(row.get("return_rate")),
            "normal_margin": nullable_number(
                row.get("lifetime_profit_margin")
            ),
            "lifetime_profit_margin": nullable_number(
                row.get("lifetime_profit_margin")
            ),
        }
    _SUPABASE_PROMOTION_CACHE.update(
        {
            "expires_at": now + max(PROMOTION_CACHE_SECONDS, 0),
            "rows": inventory,
        }
    )
    return inventory


def supabase_channeladvisor_prices():
    if not supabase_enabled():
        return {}, {}
    now = time.monotonic()
    if (
        _SUPABASE_CA_PRICE_CACHE["exact"] is not None
        and now < _SUPABASE_CA_PRICE_CACHE["expires_at"]
    ):
        return (
            _SUPABASE_CA_PRICE_CACHE["exact"],
            _SUPABASE_CA_PRICE_CACHE["canonical"],
        )

    rows = supabase_select_all(
        "channeladvisor_products",
        {
            "select": (
                "platform_sku,wooper_sku,ca_price,mapping_status"
            ),
            "ca_price": "not.is.null",
            "order": "platform_sku.asc",
        },
    )
    exact = {}
    canonical_candidates = {}
    for row in rows:
        platform_sku = str(row.get("platform_sku") or "").strip().upper()
        ca_price = nullable_number(row.get("ca_price"))
        if not platform_sku or ca_price is None:
            continue
        exact[platform_sku] = ca_price
        wooper_sku = str(row.get("wooper_sku") or "").strip().upper()
        if row.get("mapping_status") == "mapped" and wooper_sku:
            canonical_candidates.setdefault(wooper_sku, set()).add(ca_price)
    canonical = {
        sku: next(iter(prices))
        for sku, prices in canonical_candidates.items()
        if len(prices) == 1
    }
    _SUPABASE_CA_PRICE_CACHE.update(
        {
            "expires_at": now + max(PROMOTION_CACHE_SECONDS, 0),
            "exact": exact,
            "canonical": canonical,
        }
    )
    return exact, canonical


def sample_inventory():
    rows = json.loads(SAMPLE_PATH.read_text(encoding="utf-8"))["rows"]
    return {str(row["sku"]).upper(): row for row in rows}


def channeladvisor_source_path():
    if CHANNELADVISOR_PATH:
        explicit_path = Path(CHANNELADVISOR_PATH)
        if explicit_path.exists():
            return explicit_path

    downloads = Path.home() / "Downloads"
    exports = sorted(
        downloads.glob("InventoryExport_Marketplace_New_Listing_Template*.xlsx"),
        key=lambda path: path.stat().st_mtime_ns,
        reverse=True,
    )
    if exports:
        return exports[0]
    return WORKBOOK_PATH


def channeladvisor_prices():
    """Return CA prices indexed by raw SKU and the legacy naming rule."""
    if PROMOTION_DATA_SOURCE == "supabase" and supabase_enabled():
        return supabase_channeladvisor_prices()

    source_path = channeladvisor_source_path()
    if not source_path.exists():
        return {}, {}

    resolved_path = str(source_path.resolve())
    mtime_ns = source_path.stat().st_mtime_ns
    if (
        _CA_PRICE_CACHE["path"] == resolved_path
        and _CA_PRICE_CACHE["mtime_ns"] == mtime_ns
        and _CA_PRICE_CACHE["exact"] is not None
    ):
        return _CA_PRICE_CACHE["exact"], _CA_PRICE_CACHE["canonical"]

    workbook = load_workbook(source_path, data_only=True, read_only=True)
    try:
        sheet = workbook["Image"] if "Image" in workbook.sheetnames else workbook.active
        sheet.reset_dimensions()
        row_iterator = sheet.iter_rows(values_only=True)
        headers = [clean_header(value) for value in next(row_iterator, ())]
        sku_index = headers.index("inventory number")
        price_index = headers.index("buy it now price")
        exact = {}
        canonical_candidates = {}
        for values in row_iterator:
            platform_sku = str(values[sku_index] or "").strip().upper()
            price = nullable_number(values[price_index])
            if not platform_sku or price is None or price <= 0:
                continue
            exact[platform_sku] = price
            canonical = ca_price_sku(platform_sku)
            canonical_candidates.setdefault(canonical, set()).add(price)
        canonical = {
            sku: next(iter(prices))
            for sku, prices in canonical_candidates.items()
            if len(prices) == 1
        }
    finally:
        workbook.close()

    _CA_PRICE_CACHE.update(
        {
            "path": resolved_path,
            "mtime_ns": mtime_ns,
            "exact": exact,
            "canonical": canonical,
        }
    )
    return exact, canonical


def align_channeladvisor_prices(inventory_skus, mappings=None, exact_prices=None):
    """Align raw ChannelAdvisor SKUs to verified Wooper SKUs."""
    inventory_skus = {str(sku).strip().upper() for sku in inventory_skus}
    mappings = {
        str(platform_sku).strip().upper(): str(wooper_sku).strip().upper()
        for platform_sku, wooper_sku in (mappings or {}).items()
    }
    if exact_prices is None:
        exact_prices, _ = channeladvisor_prices()

    records = {}
    canonical_candidates = {}
    for platform_sku, price in exact_prices.items():
        platform_key = str(platform_sku or "").strip().upper()
        if platform_key.endswith("-ALL"):
            records[platform_key] = {
                "platform_sku": platform_key,
                "wooper_sku": None,
                "ca_price": price,
                "mapping_status": "non_existing",
                "mapping_source": "parent_sku_rule",
            }
            continue

        saved_mapping = mappings.get(platform_key)
        if saved_mapping == NON_EXISTING_SKU:
            records[platform_key] = {
                "platform_sku": platform_key,
                "wooper_sku": None,
                "ca_price": price,
                "mapping_status": "non_existing",
                "mapping_source": "saved",
            }
            continue

        wooper_sku = resolve_wooper_sku(
            platform_key,
            inventory_skus,
            mappings=mappings,
        )
        if wooper_sku is None:
            mapping_status = "unresolved"
            mapping_source = "unresolved"
        elif saved_mapping == wooper_sku:
            mapping_status = "mapped"
            mapping_source = "saved"
        elif platform_key == wooper_sku:
            mapping_status = "mapped"
            mapping_source = "exact"
        else:
            mapping_status = "mapped"
            mapping_source = "rule"

        records[platform_key] = {
            "platform_sku": platform_key,
            "wooper_sku": wooper_sku,
            "ca_price": price,
            "mapping_status": mapping_status,
            "mapping_source": mapping_source,
        }
        if wooper_sku:
            canonical_candidates.setdefault(wooper_sku, set()).add(price)

    canonical_prices = {
        sku: next(iter(prices))
        for sku, prices in canonical_candidates.items()
        if len(prices) == 1
    }
    return records, canonical_prices


def inventory_source_path():
    if WOOPER_INVENTORY_PATH:
        explicit_path = Path(WOOPER_INVENTORY_PATH)
        if explicit_path.exists():
            return explicit_path

    downloads = Path.home() / "Downloads"
    exports = sorted(
        downloads.glob("InventoryReportExport_Normal_UK_*.xlsx"),
        key=lambda path: path.stat().st_mtime_ns,
        reverse=True,
    )
    if exports:
        return exports[0]
    return WORKBOOK_PATH


def workbook_inventory():
    """Load the Wooper Inventory Report and reuse it until the workbook changes."""
    source_path = inventory_source_path()
    if not source_path.exists():
        return {}

    workbook_path = str(source_path.resolve())
    mtime_ns = source_path.stat().st_mtime_ns
    if (
        _INVENTORY_CACHE["path"] == workbook_path
        and _INVENTORY_CACHE["mtime_ns"] == mtime_ns
        and _INVENTORY_CACHE["rows"] is not None
    ):
        return _INVENTORY_CACHE["rows"]

    workbook = load_workbook(source_path, data_only=True, read_only=True)
    try:
        sheet = (
            workbook["Inventory Report"]
            if "Inventory Report" in workbook.sheetnames
            else workbook.active
        )
        sheet.reset_dimensions()
        row_iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(row_iterator, ())
        headers = [
            clean_header(str(value or "").split("/", 1)[0])
            for value in raw_headers
        ]
        field_headers = {
            "sku": "product sku",
            "main_category": "main category",
            "subcategory": "subcategory",
            "brand": "brand",
            "inventory_status": "inventory status",
            "grade": "grade level",
            "estimated_months": "estimated months to sell",
            "stock": "total inventory qty",
            "cogs": "cogs",
        }
        indexes = {
            field: headers.index(header)
            for field, header in field_headers.items()
            if header in headers
        }
        if "sku" not in indexes:
            raise ValueError("Inventory Report does not contain a Product SKU column")

        inventory = {}
        for values in row_iterator:
            sku = str(values[indexes["sku"]] or "").strip().upper()
            if not sku:
                continue
            inventory[sku] = {
                field: values[index]
                for field, index in indexes.items()
                if field != "sku"
            }
            inventory[sku]["sku"] = sku
    finally:
        workbook.close()

    _INVENTORY_CACHE.update(
        {"path": workbook_path, "mtime_ns": mtime_ns, "rows": inventory}
    )
    return inventory


def workbook_first_arrivals():
    """Match the dashboard's SKU-master then earliest-inbound date fallback."""
    source_path = WORKBOOK_PATH
    if not source_path.exists():
        return {}

    workbook_path = str(source_path.resolve())
    mtime_ns = source_path.stat().st_mtime_ns
    if (
        _FIRST_ARRIVAL_CACHE["path"] == workbook_path
        and _FIRST_ARRIVAL_CACHE["mtime_ns"] == mtime_ns
        and _FIRST_ARRIVAL_CACHE["rows"] is not None
    ):
        return _FIRST_ARRIVAL_CACHE["rows"]

    workbook = load_workbook(source_path, data_only=True, read_only=True)
    try:
        master_arrivals = {}
        if "SKU" in workbook.sheetnames:
            sheet = workbook["SKU"]
            sheet.reset_dimensions()
            row_iterator = sheet.iter_rows(values_only=True)
            headers = [clean_header(value) for value in next(row_iterator, ())]
            if "sku master" in headers and "first arrival date" in headers:
                sku_index = headers.index("sku master")
                arrival_index = headers.index("first arrival date")
                for values in row_iterator:
                    sku = str(values[sku_index] or "").strip().upper()
                    arrival = normalise_date(values[arrival_index])
                    if sku and arrival:
                        master_arrivals[sku] = arrival

        inbound_arrivals = {}
        if "Container report" in workbook.sheetnames:
            sheet = workbook["Container report"]
            sheet.reset_dimensions()
            row_iterator = sheet.iter_rows(values_only=True)
            headers = [clean_header(value) for value in next(row_iterator, ())]
            if "sku" in headers and "inbound time" in headers:
                sku_index = headers.index("sku")
                inbound_index = headers.index("inbound time")
                for values in row_iterator:
                    sku = str(values[sku_index] or "").strip().upper()
                    arrival = normalise_date(values[inbound_index])
                    if not sku or not arrival:
                        continue
                    existing = inbound_arrivals.get(sku)
                    if existing is None or arrival < existing:
                        inbound_arrivals[sku] = arrival

        arrivals = dict(inbound_arrivals)
        arrivals.update(master_arrivals)
    finally:
        workbook.close()

    _FIRST_ARRIVAL_CACHE.update(
        {"path": workbook_path, "mtime_ns": mtime_ns, "rows": arrivals}
    )
    return arrivals


def select_suggested_freight(
    suggested_freight,
    valid_qty,
    avg_actual_freight,
    sello_tools_calculation,
):
    suggested = nullable_number(suggested_freight)
    if suggested is not None:
        return suggested
    average = nullable_number(avg_actual_freight)
    if normalise_number(valid_qty) > 5 and average is not None:
        return average
    return nullable_number(sello_tools_calculation)


def workbook_suggested_freight():
    """Load the same per-unit suggested freight used by the SKU dashboard."""
    source_path = WORKBOOK_PATH
    if not source_path.exists():
        return {}

    workbook_path = str(source_path.resolve())
    mtime_ns = source_path.stat().st_mtime_ns
    if (
        _FREIGHT_CACHE["path"] == workbook_path
        and _FREIGHT_CACHE["mtime_ns"] == mtime_ns
        and _FREIGHT_CACHE["rows"] is not None
    ):
        return _FREIGHT_CACHE["rows"]

    workbook = load_workbook(source_path, data_only=True, read_only=True)
    try:
        if "Freight" not in workbook.sheetnames:
            return {}
        sheet = workbook["Freight"]
        sheet.reset_dimensions()
        row_iterator = sheet.iter_rows(values_only=True)
        headers = [clean_header(value) for value in next(row_iterator, ())]
        field_headers = {
            "sku": "sku",
            "suggested_freight": "suggested freight",
            "valid_qty": "valid qty",
            "avg_actual_freight": "avg actual freight",
            "sello_tools_calculation": "sello tools calculation",
        }
        indexes = {
            field: headers.index(header)
            for field, header in field_headers.items()
            if header in headers
        }
        if "sku" not in indexes:
            return {}

        freight = {}
        for values in row_iterator:
            sku = str(values[indexes["sku"]] or "").strip().upper()
            if not sku:
                continue
            value = select_suggested_freight(
                values[indexes["suggested_freight"]]
                if "suggested_freight" in indexes
                else None,
                values[indexes["valid_qty"]] if "valid_qty" in indexes else None,
                values[indexes["avg_actual_freight"]]
                if "avg_actual_freight" in indexes
                else None,
                values[indexes["sello_tools_calculation"]]
                if "sello_tools_calculation" in indexes
                else None,
            )
            if value is not None:
                freight[sku] = value
    finally:
        workbook.close()

    _FREIGHT_CACHE.update(
        {"path": workbook_path, "mtime_ns": mtime_ns, "rows": freight}
    )
    return freight


def aggregate_lifetime_metrics(rows):
    totals = {}
    for row in rows:
        sku = str(row.get("sku_code") or "").strip().upper()
        if not sku:
            continue
        item = totals.setdefault(
            sku,
            {
                "sold_qty": 0.0,
                "sales_amt": 0.0,
                "net_sales": 0.0,
                "return_amount": 0.0,
                "profit_incl_rn": 0.0,
            },
        )
        item["sold_qty"] += normalise_number(row.get("sku_qty"))
        item["sales_amt"] += normalise_number(row.get("sales_amt"))
        item["net_sales"] += (
            normalise_number(row.get("sales_amt"))
            + normalise_number(row.get("extra_freight"))
            - normalise_number(row.get("promo_rebate"))
        )
        item["return_amount"] += (
            normalise_number(row.get("refund_amt"))
            + normalise_number(row.get("resend_amt"))
        )
        item["profit_incl_rn"] += normalise_number(row.get("profit_incl_rn"))

    metrics = {}
    for sku, item in totals.items():
        net_sales = item["net_sales"]
        metrics[sku] = {
            **item,
            "return_rate": item["return_amount"] / net_sales if net_sales else None,
            "lifetime_profit_margin": (
                item["profit_incl_rn"] / net_sales if net_sales else None
            ),
            "normal_margin": (
                item["profit_incl_rn"] / net_sales if net_sales else None
            ),
        }
    return metrics


def workbook_lifetime_metrics():
    """Aggregate all-platform lifetime metrics using the dashboard definitions."""
    source_path = WORKBOOK_PATH
    if not source_path.exists():
        return {}

    workbook_path = str(source_path.resolve())
    mtime_ns = source_path.stat().st_mtime_ns
    if (
        _PERFORMANCE_CACHE["path"] == workbook_path
        and _PERFORMANCE_CACHE["mtime_ns"] == mtime_ns
        and _PERFORMANCE_CACHE["rows"] is not None
    ):
        return _PERFORMANCE_CACHE["rows"]

    with _PERFORMANCE_BUILD_LOCK:
        if (
            _PERFORMANCE_CACHE["path"] == workbook_path
            and _PERFORMANCE_CACHE["mtime_ns"] == mtime_ns
            and _PERFORMANCE_CACHE["rows"] is not None
        ):
            return _PERFORMANCE_CACHE["rows"]

        if LIFETIME_METRICS_PATH.exists():
            try:
                cached = json.loads(
                    LIFETIME_METRICS_PATH.read_text(encoding="utf-8")
                )
                if (
                    cached.get("source_path") == workbook_path
                    and cached.get("source_mtime_ns") == mtime_ns
                ):
                    metrics = cached.get("metrics", {})
                    _PERFORMANCE_CACHE.update(
                        {
                            "path": workbook_path,
                            "mtime_ns": mtime_ns,
                            "rows": metrics,
                        }
                    )
                    return metrics
            except (OSError, ValueError, TypeError):
                pass

        workbook = load_workbook(source_path, data_only=True, read_only=True)
        try:
            if "PowerBI" not in workbook.sheetnames:
                return {}
            sheet = workbook["PowerBI"]
            sheet.reset_dimensions()
            row_iterator = sheet.iter_rows(values_only=True)
            headers = [clean_header(value) for value in next(row_iterator, ())]
            field_headers = {
                "sku_code": "sku code",
                "sku_qty": "sku qty",
                "sales_amt": "sales amt",
                "extra_freight": "extra freight",
                "promo_rebate": "promo rebate",
                "resend_amt": "resend amt",
                "refund_amt": "refund amt",
                "profit_incl_rn": "profit incl rn",
            }
            indexes = {
                field: headers.index(header)
                for field, header in field_headers.items()
                if header in headers
            }
            if "sku_code" not in indexes:
                return {}
            metrics = aggregate_lifetime_metrics(
                {
                    field: values[index]
                    for field, index in indexes.items()
                }
                for values in row_iterator
            )
        finally:
            workbook.close()

        LIFETIME_METRICS_PATH.parent.mkdir(parents=True, exist_ok=True)
        cache_payload = {
            "source_path": workbook_path,
            "source_mtime_ns": mtime_ns,
            "metrics": metrics,
        }
        temporary_path = LIFETIME_METRICS_PATH.with_suffix(".tmp")
        temporary_path.write_text(
            json.dumps(cache_payload, ensure_ascii=True),
            encoding="utf-8",
        )
        temporary_path.replace(LIFETIME_METRICS_PATH)
        _PERFORMANCE_CACHE.update(
            {"path": workbook_path, "mtime_ns": mtime_ns, "rows": metrics}
        )
        return metrics


def wooper_inventory():
    if PROMOTION_DATA_SOURCE == "supabase" and supabase_enabled():
        return supabase_promotion_inventory()

    inventory = dict(workbook_inventory())
    for sku, first_arrival in workbook_first_arrivals().items():
        if sku in inventory:
            inventory[sku] = {
                **inventory[sku],
                "first_arrival_date": first_arrival,
            }
    for sku, suggested_freight in workbook_suggested_freight().items():
        if sku in inventory:
            inventory[sku] = {
                **inventory[sku],
                "avg_freight": suggested_freight,
            }
    for sku, metrics in workbook_lifetime_metrics().items():
        if sku in inventory:
            inventory[sku] = {**inventory[sku], **metrics}
    for sku, sample_row in sample_inventory().items():
        inventory[sku] = {**inventory.get(sku, {}), **sample_row, "sku": sku}
    return inventory


def resolve_wooper_sku(
    platform_sku,
    inventory_skus,
    mappings=None,
    explicit_wooper_sku=None,
):
    platform_key = str(platform_sku or "").strip().upper()
    inventory_skus = {str(sku).strip().upper() for sku in inventory_skus}
    mappings = mappings or {}
    candidates = [
        mappings.get(platform_key),
        explicit_wooper_sku,
        platform_key,
        ca_price_sku(platform_key),
    ]
    for candidate in candidates:
        candidate_key = str(candidate or "").strip().upper()
        if candidate_key and candidate_key in inventory_skus:
            return candidate_key
    return None


def calculate_candidate(row: dict, criteria: dict) -> dict:
    result = dict(row)
    legacy_price = nullable_number(row.get("price"))
    offer_price = nullable_number(row.get("offer_price"))
    if offer_price is None:
        offer_price = legacy_price
    ca_price = nullable_number(row.get("ca_price"))
    requested_price_source = str(criteria.get("price_source") or "ca").lower()
    valid_ca_price = ca_price is not None and ca_price > 0
    valid_offer_price = offer_price is not None and offer_price > 0
    if requested_price_source == "ca":
        price = ca_price if valid_ca_price else offer_price
        calculation_price_source = (
            "ca" if valid_ca_price else "offer" if valid_offer_price else None
        )
    else:
        price = offer_price if valid_offer_price else ca_price
        calculation_price_source = (
            "offer" if valid_offer_price else "ca" if valid_ca_price else None
        )
    price = price or 0
    cogs = normalise_number(row.get("cogs"))
    sold_qty = normalise_number(row.get("sold_qty"))
    postage = normalise_number(row.get("postage"))
    avg_freight = nullable_number(row.get("avg_freight"))
    if avg_freight is None:
        avg_freight = postage / sold_qty if sold_qty else 0

    commission = nullable_number(row.get("commission"))
    if commission is None:
        commission = normalise_number(criteria.get("default_commission"), 0.264)

    margins = criteria.get("grade_margins", {})
    target_margin = normalise_number(margins.get(grade_key(row.get("grade"))), 0.12)
    max_discount = normalise_number(criteria.get("max_discount"), 0.25)
    wms_rate = normalise_number(criteria.get("wms_rate"), 0.08)
    vat_rate = normalise_number(criteria.get("vat_rate"), 0.20)
    input_price_includes_vat = criteria.get("input_price_includes_vat", True)
    export_price_includes_vat = criteria.get("export_price_includes_vat", True)
    vat_multiplier = 1 + vat_rate
    calculation_price_includes_vat = (
        True if calculation_price_source == "ca" else input_price_includes_vat
    )
    price_including_vat = (
        price if calculation_price_includes_vat else price * vat_multiplier
    )

    denominator = 1 - commission - wms_rate - target_margin
    if price_including_vat <= 0 or denominator <= 0:
        suggested_discount = 0
        calculation_error = "Price or margin settings make reverse pricing impossible"
    else:
        suggested_discount = (
            1
            - (
                ((cogs + avg_freight) / denominator)
                * vat_multiplier
                / price_including_vat
            )
        )
        suggested_discount = min(suggested_discount, max_discount)
        discount_interval = normalise_number(
            criteria.get("discount_interval"),
            0.05,
        )
        if (
            criteria.get("use_discount_interval", False)
            and discount_interval > 0
            and suggested_discount > 0
        ):
            suggested_discount = (
                math.floor((suggested_discount + 1e-12) / discount_interval)
                * discount_interval
            )
        calculation_error = ""

    override_discount = nullable_number(row.get("override_discount"))
    final_discount = (
        override_discount if override_discount is not None else suggested_discount
    )
    raw_promo_price = price_including_vat * (1 - final_discount)
    if criteria.get("rounding", True):
        promo_price_including_vat = excel_mround(raw_promo_price, 0.1)
        if final_discount < max_discount:
            promo_price_including_vat -= 0.05
    else:
        promo_price_including_vat = raw_promo_price
    promo_price_including_vat = max(promo_price_including_vat, 0)

    ex_vat_price = (
        promo_price_including_vat / vat_multiplier
        if vat_multiplier
        else promo_price_including_vat
    )
    promo_price = (
        promo_price_including_vat
        if export_price_includes_vat
        else ex_vat_price
    )
    wms_fee = ex_vat_price * wms_rate
    profit = (
        ex_vat_price * (1 - commission)
        - cogs
        - avg_freight
        - wms_fee
    )
    promo_margin = profit / ex_vat_price if ex_vat_price else None

    reasons = []
    warnings = []
    if requested_price_source == "offer" and calculation_price_source == "ca":
        warnings.append("Offer price unavailable; CA price used")
    if requested_price_source == "ca" and calculation_price_source == "offer":
        warnings.append("CA price unavailable; offer price used")
    if row.get("commission_requires_review"):
        warnings.append("Commission rate needs confirmation")
    grade = normalise_number(row.get("grade"))
    stock = normalise_number(row.get("stock"))
    estimated_months = normalise_number(row.get("estimated_months"))
    return_rate = normalise_number(row.get("return_rate"))
    minimum_discount = normalise_number(criteria.get("min_discount"), 0)

    if grade < normalise_number(criteria.get("min_grade"), 0):
        reasons.append("Grade below threshold")
    if stock < normalise_number(criteria.get("min_stock"), 0):
        reasons.append("Stock below threshold")
    if estimated_months < normalise_number(criteria.get("min_months"), 0):
        reasons.append("Saleable months below threshold")
    if suggested_discount < minimum_discount:
        reasons.append("Available discount below minimum")
    if final_discount < 0:
        reasons.append("Target margin requires a price increase")

    categories = [
        clean_header(value)
        for value in criteria.get("categories", [])
        if clean_header(value)
    ]
    if categories:
        candidate_category = clean_header(row.get("subcategory"))
        if not any(category in candidate_category for category in categories):
            reasons.append("Outside selected categories")

    def selected_values(plural_key: str, singular_key: str) -> list[str]:
        raw_values = criteria.get(plural_key)
        if raw_values is None:
            raw_values = criteria.get(singular_key)
        if not isinstance(raw_values, (list, tuple, set)):
            raw_values = [raw_values]
        return [clean_header(value) for value in raw_values if clean_header(value)]

    main_categories = selected_values("main_categories", "main_category")
    if (
        main_categories
        and clean_header(row.get("main_category")) not in main_categories
    ):
        reasons.append("Outside selected main category")
    subcategories = selected_values("subcategories", "subcategory")
    if subcategories and clean_header(row.get("subcategory")) not in subcategories:
        reasons.append("Outside selected subcategory")
    brands = selected_values("brands", "brand")
    if brands and clean_header(row.get("brand")) not in brands:
        reasons.append("Outside selected brand")

    cutoff = normalise_date(criteria.get("exclude_first_arrival_on_or_after"))
    first_arrival = normalise_date(row.get("first_arrival_date"))
    if cutoff and first_arrival and first_arrival >= cutoff:
        reasons.append("First arrival is inside the excluded new-product period")

    protected_skus = {
        str(value or "").strip().upper()
        for value in criteria.get("protected_skus", [])
        if str(value or "").strip()
    }
    protection_list_excluded = (
        bool(criteria.get("exclude_current_protection_list"))
        and str(row.get("sku") or "").strip().upper() in protected_skus
    )
    if protection_list_excluded:
        reasons.append("SKU is in the current protection period")

    return_review_threshold = normalise_number(
        criteria.get("max_return_rate"),
        0.06,
    )
    return_rate_review = return_rate >= return_review_threshold
    if return_rate_review:
        warnings.append(
            "Return rate is at or above "
            f"{return_review_threshold * 100:g}% review threshold"
        )

    lifetime_margin = nullable_number(row.get("lifetime_profit_margin"))
    lifetime_margin_gap = (
        lifetime_margin - promo_margin
        if lifetime_margin is not None and promo_margin is not None
        else None
    )
    margin_gap_review = (
        lifetime_margin_gap is not None
        and lifetime_margin_gap > 0.05 + 1e-12
    )
    if margin_gap_review:
        warnings.append(
            "Lifetime margin exceeds promo margin by more than 5 points"
        )

    if promo_margin is not None and promo_margin + 0.0001 < target_margin:
        if override_discount is not None:
            warnings.append("Override discount leaves margin below target")
        elif criteria.get("rounding", True):
            warnings.append("Rounding leaves margin slightly below target")
        else:
            warnings.append("Final discount leaves margin below target")
    if calculation_error:
        reasons.append(calculation_error)

    result.update(
        {
            "avg_freight": avg_freight,
            "commission": commission,
            "offer_price": offer_price,
            "ca_price": ca_price,
            "price": price,
            "price_including_vat": price_including_vat,
            "calculation_price_source": calculation_price_source,
            "requested_price_source": requested_price_source,
            "target_margin": target_margin,
            "suggested_discount": suggested_discount,
            "final_discount": final_discount,
            "use_discount_interval": criteria.get(
                "use_discount_interval",
                False,
            ),
            "discount_interval": normalise_number(
                criteria.get("discount_interval"),
                0.05,
            ),
            "promo_price": promo_price,
            "promo_price_including_vat": promo_price_including_vat,
            "promo_price_excluding_vat": ex_vat_price,
            "input_price_includes_vat": input_price_includes_vat,
            "export_price_includes_vat": export_price_includes_vat,
            "promo_profit": profit,
            "promo_margin": promo_margin,
            "return_rate_review": return_rate_review,
            "lifetime_margin_gap": lifetime_margin_gap,
            "margin_gap_review": margin_gap_review,
            "protection_list_excluded": protection_list_excluded,
            "wms_fee": wms_fee,
            "eligible": not reasons,
            "reasons": reasons,
            "warnings": warnings,
        }
    )
    return result


def rows_from_csv(data: bytes):
    text = data.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def rows_from_workbook(data: bytes):
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheet = workbook.active
    sheet.reset_dimensions()
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator, None)
    if not headers:
        return []
    return [
        {str(headers[index] or ""): value for index, value in enumerate(values)}
        for values in iterator
        if any(value not in (None, "") for value in values)
    ]


def normalise_commission_rows(rows, default_platform=""):
    platform_headers = {"platform", "platform name", "marketplace"}
    category_headers = {"category", "product category", "commission category"}
    sku_headers = {
        "sku",
        "sku code",
        "wooper sku",
        "core sku",
        "inventory number",
    }
    commission_headers = ALIASES["commission"] | {"rate"}
    normalised = {}
    conflicts = []

    for source_row in rows:
        cleaned = {
            clean_header(header): value
            for header, value in source_row.items()
        }

        def first_value(headers):
            return next(
                (
                    cleaned[header]
                    for header in headers
                    if header in cleaned and cleaned[header] not in (None, "")
                ),
                None,
            )

        platform = str(first_value(platform_headers) or default_platform).strip()
        sku = str(first_value(sku_headers) or "").strip().upper()
        category = str(first_value(category_headers) or "").strip()
        commission = nullable_number(first_value(commission_headers))
        if commission is not None and commission > 1:
            commission /= 100
        if not platform or not sku or commission is None:
            continue
        if commission < 0 or commission > 1:
            continue

        canonical_platform = next(
            (
                name
                for name in DEFAULT_COMMISSIONS
                if clean_header(name) == clean_header(platform)
            ),
            platform,
        )
        key = (clean_header(canonical_platform), sku)
        existing = normalised.get(key)
        if existing and not math.isclose(
            existing["commission"],
            commission,
            abs_tol=1e-12,
        ):
            conflicts.append(
                {
                    "platform": canonical_platform,
                    "sku": sku,
                    "rates": [existing["commission"], commission],
                }
            )
            continue
        normalised[key] = {
            "platform": canonical_platform,
            "category": category,
            "sku": sku,
            "commission": commission,
        }

    return list(normalised.values()), conflicts


def commission_rows_from_workbook(data: bytes):
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows = []
    conflicts = []
    for sheet in workbook.worksheets:
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            continue
        source_rows = [
            {
                str(headers[index] or ""): value
                for index, value in enumerate(values)
            }
            for values in iterator
            if any(value not in (None, "") for value in values)
        ]
        sheet_rows, sheet_conflicts = normalise_commission_rows(
            source_rows,
            default_platform=sheet.title,
        )
        rows.extend(sheet_rows)
        conflicts.extend(sheet_conflicts)
    return rows, conflicts


def map_imported_rows(
    rows,
    manual_mappings=None,
    inventory=None,
    exact_ca_prices=None,
    platform="",
):
    alias_lookup = {
        alias: field for field, aliases in ALIASES.items() for alias in aliases
    }
    mapped = []
    if manual_mappings is None:
        manual_mappings = combined_platform_mappings(platform)
    if inventory is None:
        inventory = wooper_inventory()
    if exact_ca_prices is None:
        exact_ca_prices, _ = channeladvisor_prices()
    _, aligned_ca_prices = align_channeladvisor_prices(
        inventory,
        mappings=manual_mappings,
        exact_prices=exact_ca_prices,
    )
    for source_row in rows:
        target = {}
        for header, value in source_row.items():
            field = alias_lookup.get(clean_header(header))
            if field and field not in target:
                target[field] = value
        platform_sku = str(
            target.get("platform_sku")
            or target.get("sku")
            or target.get("wooper_sku")
            or ""
        ).strip()
        if not platform_sku:
            continue
        target["platform_sku"] = platform_sku
        platform_key = platform_sku.upper()
        if manual_mappings.get(platform_key) == NON_EXISTING_SKU:
            continue
        automatic_sku = ca_price_sku(platform_sku)
        canonical = resolve_wooper_sku(
            platform_sku,
            inventory,
            mappings=manual_mappings,
            explicit_wooper_sku=target.get("wooper_sku"),
        )
        matched = canonical is not None
        imported_ca_price = nullable_number(target.get("ca_price"))
        target["ca_price"] = (
            imported_ca_price
            if imported_ca_price is not None
            else exact_ca_prices.get(platform_key)
            or aligned_ca_prices.get(canonical or automatic_sku)
        )
        target["suggested_sku"] = canonical or ""
        target["sku"] = canonical or ""
        target["mapping_status"] = "mapped" if matched else "unresolved"
        if matched:
            platform_values = dict(target)
            target = dict(inventory[canonical])
            target.pop("price", None)
            target.pop("offer_price", None)
            target.pop("ca_price", None)
            target.update(platform_values)
            target["sku"] = canonical
        mapped.append(target)
    return mapped


def source_value(row, names, default=None):
    cleaned = {clean_header(header): value for header, value in row.items()}
    for name in names:
        value = cleaned.get(clean_header(name))
        if value not in (None, ""):
            return value
    return default


def write_json_file(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temporary.replace(path)


def read_json_file(path, default):
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def capture_tesco_offers(source_rows, mapped_rows):
    mapped_by_sku = {
        str(row.get("platform_sku") or "").strip().upper(): row
        for row in mapped_rows
    }
    offers = []
    for row in source_rows:
        tesco_sku = str(
            source_value(
                row,
                ("Platform SKU", "SKU", "Merchant SKU", "Seller SKU"),
                "",
            )
        ).strip().upper()
        if not tesco_sku:
            continue
        mapped = mapped_by_sku.get(tesco_sku, {})
        price = source_value(
            row,
            ("Buy It Now Price", "Buy Now Price", "CA Price", "Normal Price", "Price"),
        )
        offers.append(
            {
                "tesco_sku": tesco_sku,
                "product_id": str(
                    source_value(row, ("Product ID", "Product-ID", "Offer ID"), "")
                ).strip(),
                "wooper_sku": str(mapped.get("sku") or "").strip().upper(),
                "ca_price": nullable_number(mapped.get("ca_price"))
                or nullable_number(price),
                "mapping_status": mapped.get("mapping_status", "unresolved"),
            }
        )
    captured_at = datetime.now(ZoneInfo("Australia/Sydney")).isoformat()
    if supabase_enabled():
        stored_offers = [dict(offer, captured_at=captured_at) for offer in offers]
        if stored_offers:
            supabase_request(
                "POST",
                "promotion_tesco_offers",
                rows=stored_offers,
                params={"on_conflict": "tesco_sku"},
                prefer="resolution=merge-duplicates,return=minimal",
            )
        supabase_request(
            "DELETE",
            "promotion_tesco_offers",
            params={"captured_at": f"lt.{captured_at}"},
        )
    else:
        write_json_file(
            TESCO_OFFERS_PATH,
            {"captured_at": captured_at, "offers": offers},
        )
    return offers


def latest_tesco_offers():
    if supabase_enabled():
        return supabase_select_all(
            "promotion_tesco_offers",
            {
                "select": (
                    "tesco_sku,product_id,wooper_sku,ca_price,"
                    "mapping_status,captured_at"
                ),
                "order": "tesco_sku.asc",
            },
        )
    return read_json_file(TESCO_OFFERS_PATH, {"offers": []}).get("offers", [])


def latest_tesco_offers_captured_at(offers=None):
    offers = offers if offers is not None else latest_tesco_offers()
    if offers and supabase_enabled():
        return max(
            (str(row.get("captured_at") or "") for row in offers),
            default="",
        ) or None
    return read_json_file(TESCO_OFFERS_PATH, {}).get("captured_at")


def tesco_candidates_from_rows(rows, inventory=None):
    inventory = inventory or {}
    candidates = []
    seen_tesco_skus = set()
    for row in rows:
        approval = str(source_value(row, ("Approval",), "")).strip().lower()
        if approval not in {"yes", "y", "true", "1"}:
            continue
        platform = str(source_value(row, ("Platform",), "")).strip()
        if platform and clean_header(platform) != "tesco":
            continue
        tesco_sku = str(
            source_value(row, ("Platform SKU", "Tesco SKU", "SKU"), "")
        ).strip().upper()
        wooper_sku = str(
            source_value(row, ("SKU", "Wooper SKU", "Core SKU"), "")
        ).strip().upper()
        if not tesco_sku or tesco_sku in seen_tesco_skus:
            continue
        seen_tesco_skus.add(tesco_sku)
        performance = inventory.get(wooper_sku, {})
        candidates.append(
            {
                "tesco_sku": tesco_sku,
                "wooper_sku": wooper_sku,
                "product_id": str(source_value(row, ("Product ID",), "")).strip(),
                "ca_price": nullable_number(
                    source_value(
                        row,
                        (
                            "CA Price - Normal Price (VAT Included)",
                            "CA Price",
                            "Normal Price",
                        ),
                    )
                ),
                "promo_price": nullable_number(
                    source_value(
                        row,
                        (
                            "Promotion Price (VAT Included)",
                            "Promotion Price (VAT Excluded)",
                            "Promotion Price",
                        ),
                    )
                ),
                "discount": nullable_number(source_value(row, ("Final Discount",))),
                "soh": normalise_number(source_value(row, ("SOH", "Stock")), 0),
                "grade": nullable_number(source_value(row, ("Grade", "Grade Level")))
                if source_value(row, ("Grade", "Grade Level")) not in (None, "")
                else nullable_number(performance.get("grade")),
                "months": nullable_number(source_value(row, ("Months",)))
                if source_value(row, ("Months",)) not in (None, "")
                else nullable_number(performance.get("estimated_months")),
                "lifetime_margin": nullable_number(
                    source_value(row, ("Lifetime Profit Margin (After Returns)",))
                ),
                "return_rate": nullable_number(
                    source_value(row, ("Return Rate (All Platforms)",))
                ),
            }
        )
    return candidates


def tesco_catalogue_from_workbook(data, inventory=None, mappings=None):
    inventory = inventory or {}
    mappings = mappings or {}
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    best_rows = []
    for sheet in workbook.worksheets:
        sheet.reset_dimensions()
        values = sheet.iter_rows(values_only=True)
        header = None
        for _ in range(12):
            possible = next(values, None)
            if possible is None:
                break
            names = {clean_header(value) for value in possible if value not in (None, "")}
            if "sku" in names and ("title" in names or "product title" in names):
                header = [str(value or "") for value in possible]
                break
        if not header:
            continue
        rows = []
        for item in values:
            source = {
                header[index]: value
                for index, value in enumerate(item[: len(header)])
            }
            tesco_sku = str(source_value(source, ("SKU", "Tesco SKU"), "")).strip().upper()
            if not tesco_sku or tesco_sku in {"SKU", "TESCO SKU"}:
                continue
            image = source_value(
                source,
                ("Image 1", "Image URL 1", "Main Image", "Image URL", "Image"),
                "",
            )
            rows.append(
                {
                    "tesco_sku": tesco_sku,
                    "wooper_sku": resolve_wooper_sku(
                        tesco_sku,
                        inventory,
                        mappings=mappings,
                    ) or "",
                    "category_path": str(
                        source_value(source, ("Category Code", "Cat Path", "Category Path"), "")
                    ).strip(),
                    "title": str(source_value(source, ("Title", "Product Title"), "")).strip(),
                    "brand": str(source_value(source, ("Brand",), "")).strip(),
                    "barcode": str(
                        source_value(source, ("Barcode", "EAN", "EAN / GTIN"), "")
                    ).strip(),
                    "image_url": str(image or "").strip(),
                }
            )
        if len(rows) > len(best_rows):
            best_rows = rows
    return best_rows


def capture_tesco_catalogue(rows):
    captured_at = datetime.now(ZoneInfo("Australia/Sydney")).isoformat()
    catalogue = []
    seen = set()
    for source in rows:
        tesco_sku = str(source.get("tesco_sku") or "").strip().upper()
        if not tesco_sku or tesco_sku in seen:
            continue
        seen.add(tesco_sku)
        catalogue.append(
            {
                "tesco_sku": tesco_sku,
                "wooper_sku": str(source.get("wooper_sku") or "").strip().upper(),
                "barcode": normalise_barcode(source.get("barcode")),
                "title": str(source.get("title") or "").strip(),
                "brand": str(source.get("brand") or "").strip(),
                "category_path": str(source.get("category_path") or "").strip(),
                "image_url": str(source.get("image_url") or "").strip(),
                "captured_at": captured_at,
            }
        )
    if supabase_enabled():
        for offset in range(0, len(catalogue), 500):
            supabase_request(
                "POST",
                "promotion_tesco_catalogue",
                rows=catalogue[offset : offset + 500],
                params={"on_conflict": "tesco_sku"},
                prefer="resolution=merge-duplicates,return=minimal",
            )
        supabase_request(
            "DELETE",
            "promotion_tesco_catalogue",
            params={"captured_at": f"lt.{captured_at}"},
        )
    else:
        write_json_file(
            TESCO_CATALOGUE_PATH,
            {"captured_at": captured_at, "catalogue": catalogue},
        )
    return catalogue


def latest_tesco_catalogue():
    if supabase_enabled():
        return supabase_select_all(
            "promotion_tesco_catalogue",
            {
                "select": (
                    "tesco_sku,wooper_sku,barcode,title,brand,category_path,"
                    "image_url,captured_at"
                ),
                "order": "tesco_sku.asc",
            },
        )
    return read_json_file(TESCO_CATALOGUE_PATH, {"catalogue": []}).get(
        "catalogue", []
    )


def normalise_barcode(value):
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def normalise_excel_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text:
        return ""
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Could not read nomination date: {text}")


def workbook_table_rows(sheet, required_headers, scan_rows=15):
    sheet.reset_dimensions()
    values = sheet.iter_rows(values_only=True)
    for _ in range(scan_rows):
        possible = next(values, None)
        if possible is None:
            break
        headers = [str(value or "") for value in possible]
        cleaned = {clean_header(value) for value in headers if value not in (None, "")}
        if all(any(name in cleaned for name in choices) for choices in required_headers):
            return [
                {headers[index]: value for index, value in enumerate(row[: len(headers)])}
                for row in values
            ]
    return []


def tesco_external_nomination_from_workbook(data, catalogue=None):
    catalogue = catalogue if catalogue is not None else latest_tesco_catalogue()
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    official_rows = []
    for sheet in workbook.worksheets:
        rows = workbook_table_rows(
            sheet,
            (
                {"event", "event line"},
                {"start date"},
                {"end date"},
                {"barcode", "barcode/ean", "barcode / ean"},
            ),
        )
        if rows:
            official_rows = rows
            break
    if not official_rows:
        raise ValueError("No Tesco nomination rows were found in this workbook.")

    internal_rows = []
    if "Internal Record" in workbook.sheetnames:
        internal_rows = workbook_table_rows(
            workbook["Internal Record"],
            ({"platform sku"}, {"wooper sku"}),
        )
    else:
        # Tesco's standard nomination template always uses row 3 as an example.
        official_rows = official_rows[1:]

    by_sku = {
        str(row.get("tesco_sku") or "").strip().upper(): row
        for row in catalogue
        if str(row.get("tesco_sku") or "").strip()
    }
    by_barcode = {}
    for row in catalogue:
        barcode = normalise_barcode(row.get("barcode"))
        if barcode:
            by_barcode.setdefault(barcode, []).append(row)

    grouped = {}
    official_index = 0
    for source in official_rows:
        event_name = str(source_value(source, ("Event", "Event line"), "")).strip()
        barcode = normalise_barcode(
            source_value(source, ("Barcode / EAN", "Barcode/EAN", "Barcode", "EAN"), "")
        )
        title = str(source_value(source, ("Product Title", "Title"), "")).strip()
        if not event_name and not barcode and not title:
            continue
        start_date = normalise_excel_date(source_value(source, ("Start Date",), ""))
        end_date = normalise_excel_date(source_value(source, ("End Date",), ""))
        if not event_name or not start_date or not end_date:
            raise ValueError("Every nomination row needs an event name, start date, and end date.")
        if start_date > end_date:
            raise ValueError(f"The end date is before the start date for {event_name}.")

        internal = internal_rows[official_index] if official_index < len(internal_rows) else {}
        official_index += 1
        direct_sku = str(source_value(internal, ("Platform SKU",), "")).strip().upper()
        direct_wooper = str(source_value(internal, ("Wooper SKU",), "")).strip().upper()
        matches = by_barcode.get(barcode, []) if barcode else []
        if direct_sku:
            matched = by_sku.get(direct_sku, {})
            tesco_sku = direct_sku
            wooper_sku = str(matched.get("wooper_sku") or direct_wooper).strip().upper()
            match_status = "matched"
            match_source = "internal_record"
        elif len(matches) == 1:
            tesco_sku = str(matches[0].get("tesco_sku") or "").strip().upper()
            wooper_sku = str(matches[0].get("wooper_sku") or "").strip().upper()
            match_status = "matched"
            match_source = "barcode"
        else:
            tesco_sku = ""
            wooper_sku = ""
            match_status = "ambiguous" if len(matches) > 1 else "unmatched"
            match_source = ""
        key = (event_name, start_date, end_date)
        grouped.setdefault(
            key,
            {
                "event_name": event_name,
                "start_date": start_date,
                "end_date": end_date,
                "rows": [],
            },
        )["rows"].append(
            {
                "tesco_sku": tesco_sku,
                "wooper_sku": wooper_sku,
                "barcode": barcode,
                "title": title,
                "match_status": match_status,
                "match_source": match_source,
                "suggestions": [
                    {
                        "tesco_sku": str(item.get("tesco_sku") or "").strip().upper(),
                        "wooper_sku": str(item.get("wooper_sku") or "").strip().upper(),
                        "title": str(item.get("title") or "").strip(),
                    }
                    for item in matches
                ],
            }
        )
    events = list(grouped.values())
    if not events:
        raise ValueError("No completed Tesco nomination rows were found.")
    return {
        "events": events,
        "catalogue_options": [
            {
                "tesco_sku": sku,
                "wooper_sku": str(row.get("wooper_sku") or "").strip().upper(),
                "title": str(row.get("title") or "").strip(),
            }
            for sku, row in sorted(by_sku.items())
        ],
    }


def tesco_events():
    if supabase_enabled():
        events = supabase_select_all(
            "promotion_tesco_events",
            {
                "select": (
                    "id,event_name,start_date,end_date,discount_end_date_mirakl,"
                    "created_at,source,rows"
                ),
                "order": "created_at.asc",
            },
        )
        if events:
            return events
        local_events = read_json_file(TESCO_EVENTS_PATH, {"events": []}).get(
            "events", []
        )
        if not local_events:
            return []
        seed_rows = []
        for event in local_events:
            source = event.get("source")
            if not isinstance(source, dict):
                source = {"file": str(source or "Promotion Master.xlsx")}
            seed_rows.append(
                {
                    "id": str(
                        uuid.uuid5(
                            uuid.NAMESPACE_URL,
                            f"tesco-promotion-event:{event.get('id')}",
                        )
                    ),
                    "event_name": str(event.get("event_name") or "Saved event"),
                    "start_date": str(event.get("start_date") or ""),
                    "end_date": str(event.get("end_date") or ""),
                    "discount_end_date_mirakl": str(
                        event.get("discount_end_date_mirakl") or ""
                    ).strip() or None,
                    "created_at": str(event.get("created_at") or datetime.now(ZoneInfo("Australia/Sydney")).isoformat()),
                    "source": source,
                    "rows": event.get("rows") or [],
                }
            )
        return supabase_request(
            "POST",
            "promotion_tesco_events",
            rows=seed_rows,
            params={"on_conflict": "id"},
            prefer="resolution=ignore-duplicates,return=representation",
        )
    return read_json_file(TESCO_EVENTS_PATH, {"events": []}).get("events", [])


def dates_overlap(start_a, end_a, start_b, end_b):
    return bool(start_a and end_a and start_b and end_b and start_a <= end_b and start_b <= end_a)


def tesco_conflicts(start_date, end_date):
    conflicts = {}
    for event in tesco_events():
        if not dates_overlap(
            str(event.get("start_date") or ""),
            str(event.get("end_date") or ""),
            start_date,
            end_date,
        ):
            continue
        for row in event.get("rows", []):
            sku = str(row.get("tesco_sku") or "").strip().upper()
            if sku:
                conflicts.setdefault(sku, []).append(event.get("event_name") or "Saved event")
    return conflicts


def build_tesco_nomination_workbook(payload):
    rows = payload.get("rows") or []
    if not rows:
        raise ValueError("Select at least one Tesco SKU.")
    event_name = str(payload.get("event_name") or "").strip()
    start_date = str(payload.get("start_date") or "").strip()
    end_date = str(payload.get("end_date") or "").strip()
    if not event_name or not start_date or not end_date:
        raise ValueError("Event name, start date, and end date are required.")
    total_limit = int(normalise_number(payload.get("total_limit"), 0))
    if total_limit and len(rows) > total_limit:
        raise ValueError("The total Tesco SKU limit has been exceeded.")
    category_limits = {
        str(item.get("name") or "").strip(): int(normalise_number(item.get("limit"), 0))
        for item in payload.get("category_limits", [])
        if str(item.get("name") or "").strip()
    }
    for category, limit in category_limits.items():
        category_count = sum(row.get("event_category") == category for row in rows)
        if limit and category_count > limit:
            raise ValueError(f"The {category} SKU limit has been exceeded.")
    conflicts = tesco_conflicts(start_date, end_date)
    blocked = [row.get("tesco_sku") for row in rows if str(row.get("tesco_sku") or "").upper() in conflicts]
    if blocked:
        raise ValueError("Overlapping Tesco SKU(s): " + ", ".join(blocked))
    for row in rows:
        if not str(row.get("event_category") or "").strip():
            raise ValueError(f"Select an event category for {row.get('tesco_sku')}.")
        if row.get("level_1") not in TESCO_LEVEL_1:
            raise ValueError(f"Select a valid Level 1 group for {row.get('tesco_sku')}.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    headers = [
        "Seller Name",
        "Brand",
        "Event",
        "Level 1",
        "Level 2",
        "Start Date",
        "End Date",
        "Barcode / EAN",
        "TPNB",
        "TPNC",
        "Product Title",
        "Was",
        "Now",
        "Discount £",
        "Discount %",
        "Hero from Seller",
        "Stock",
    ]
    sheet.append(
        [
            "Seller Input", "", "", "", "", "", "", "",
            "Calcs - do not overtype", "", "Seller Input", "", "",
            "Calcs - do not overtype", "", "Seller Input", "",
        ]
    )
    sheet.merge_cells("A1:H1")
    sheet.merge_cells("I1:J1")
    sheet.merge_cells("N1:O1")
    sheet.merge_cells("P1:Q1")
    sheet.append(headers)
    group_fill = PatternFill("solid", fgColor="D9EAD3")
    calc_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.fill = calc_fill if cell.column in {9, 10, 14, 15} else group_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill("solid", fgColor="1C2B27")
    for cell in sheet[2]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    seller_name = str(payload.get("seller_name") or "Traderight Group").strip()
    for index, row in enumerate(rows, start=3):
        sheet.append(
            [
                seller_name,
                row.get("brand", ""),
                event_name,
                row.get("level_1", ""),
                row.get("category_path", ""),
                datetime.fromisoformat(start_date),
                datetime.fromisoformat(end_date),
                row.get("barcode", ""),
                "",
                "",
                row.get("title", ""),
                normalise_number(row.get("ca_price"), 0),
                normalise_number(row.get("promo_price"), 0),
                f"=L{index}-M{index}",
                f"=IFERROR((L{index}-M{index})/L{index},0)",
                "Exclusive Pricing",
                int(normalise_number(row.get("soh"), 0)) + 100,
            ]
        )
        sheet.cell(index, 6).number_format = "dd/mm/yyyy"
        sheet.cell(index, 7).number_format = "dd/mm/yyyy"
        for column in (12, 13, 14):
            sheet.cell(index, column).number_format = "£0.00"
        sheet.cell(index, 15).number_format = "0.0%"
    internal = workbook.create_sheet("Internal Record")
    internal_headers = [
        "Platform SKU",
        "Wooper SKU",
        "Original Price (CA Price, VAT Included)",
        "Discounted Price",
        "Discount %",
    ]
    internal.append(internal_headers)
    for cell in internal[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        internal.append(
            [
                row.get("tesco_sku", ""),
                row.get("wooper_sku", ""),
                normalise_number(row.get("ca_price"), 0),
                normalise_number(row.get("promo_price"), 0),
                normalise_number(row.get("discount"), 0),
            ]
        )
    for row_number in range(2, len(rows) + 2):
        internal.cell(row_number, 3).number_format = "£0.00"
        internal.cell(row_number, 4).number_format = "£0.00"
        internal.cell(row_number, 5).number_format = "0.0%"
    for column, width in {"A": 24, "B": 24, "C": 24, "D": 18, "E": 14}.items():
        internal.column_dimensions[column].width = width
    internal.freeze_panes = "A2"
    internal.auto_filter.ref = f"A1:E{len(rows) + 1}"

    levels = workbook.create_sheet("Sheet2")
    for index, level in enumerate(TESCO_LEVEL_1, start=1):
        levels.cell(index, 1, level)
    validation = DataValidation(type="list", formula1="=Sheet2!$A$1:$A$10")
    sheet.add_data_validation(validation)
    validation.add("D3:D70")
    widths = [20, 18, 24, 20, 28, 13, 13, 18, 11, 11, 48, 12, 12, 14, 12, 20, 11]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:Q{len(rows) + 2}"
    sheet.conditional_formatting.add(
        f"A3:Q{len(rows) + 2}",
        FormulaRule(formula=["$L3<=$M3"], fill=PatternFill("solid", fgColor="FCE8E6")),
    )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def save_tesco_event(payload, workbook_bytes=None, source=None):
    event_id = str(uuid.uuid4())
    created_at = datetime.now(ZoneInfo("Australia/Sydney")).isoformat()
    record = {
        "id": event_id,
        "event_name": str(payload.get("event_name") or "").strip(),
        "start_date": str(payload.get("start_date") or "").strip(),
        "end_date": str(payload.get("end_date") or "").strip(),
        "discount_end_date_mirakl": str(
            payload.get("discount_end_date_mirakl") or ""
        ).strip() or None,
        "created_at": created_at,
        "source": source or {
            "candidate_file": str(payload.get("candidate_file") or "").strip(),
            "catalogue_file": str(payload.get("catalogue_file") or "").strip(),
        },
        "rows": payload.get("rows") or [],
    }
    if supabase_enabled():
        saved = supabase_request(
            "POST",
            "promotion_tesco_events",
            rows=[record],
            prefer="return=representation",
        )
        if not saved:
            raise RuntimeError("Supabase did not return the saved Tesco event.")
        return saved[0]
    events = tesco_events()
    events.append(record)
    write_json_file(TESCO_EVENTS_PATH, {"events": events})
    if workbook_bytes is not None:
        TESCO_NOMINATIONS_DIR.mkdir(parents=True, exist_ok=True)
        (TESCO_NOMINATIONS_DIR / f"{event_id}.xlsx").write_bytes(workbook_bytes)
    return record


def update_tesco_event(event_id, payload):
    event_id = validate_worktable_id(event_id) if supabase_enabled() else str(event_id)
    value = str(payload.get("discount_end_date_mirakl") or "").strip()
    if len(value) > 120:
        raise ValueError("The Mirakl discount end date must be 120 characters or less.")
    update = {"discount_end_date_mirakl": value or None}
    if supabase_enabled():
        rows = supabase_request(
            "PATCH",
            "promotion_tesco_events",
            rows=update,
            params={"id": f"eq.{event_id}"},
            prefer="return=representation",
        )
        return rows[0] if rows else None
    events = tesco_events()
    record = next((event for event in events if str(event.get("id")) == event_id), None)
    if record is None:
        return None
    record.update(update)
    write_json_file(TESCO_EVENTS_PATH, {"events": events})
    return record


def import_external_tesco_events(payload, source_filename=""):
    events = payload.get("events")
    if not isinstance(events, list) or not events:
        raise ValueError("No Tesco event was supplied.")
    catalogue = latest_tesco_catalogue()
    by_sku = {
        str(row.get("tesco_sku") or "").strip().upper(): row
        for row in catalogue
        if str(row.get("tesco_sku") or "").strip()
    }
    existing = {
        (
            str(event.get("event_name") or "").strip().casefold(),
            str(event.get("start_date") or ""),
            str(event.get("end_date") or ""),
        )
        for event in tesco_events()
    }
    prepared = []
    for event in events:
        event_name = str(event.get("event_name") or "").strip()
        start_date = normalise_excel_date(event.get("start_date"))
        end_date = normalise_excel_date(event.get("end_date"))
        if not event_name or not start_date or not end_date:
            raise ValueError("Event name, start date, and end date are required.")
        if len(event_name) > 160:
            raise ValueError("The event name must be 160 characters or less.")
        if start_date > end_date:
            raise ValueError(f"The end date is before the start date for {event_name}.")
        identity = (event_name.casefold(), start_date, end_date)
        if identity in existing:
            raise ValueError(
                f"{event_name} ({start_date} to {end_date}) is already saved. "
                "Remove the old record first if it needs replacing."
            )
        rows = []
        seen = set()
        for source_row in event.get("rows") or []:
            tesco_sku = str(source_row.get("tesco_sku") or "").strip().upper()
            if not tesco_sku:
                raise ValueError("Choose a Tesco SKU for every imported nomination row.")
            if tesco_sku in seen:
                continue
            seen.add(tesco_sku)
            catalogue_row = by_sku.get(tesco_sku, {})
            rows.append(
                {
                    "tesco_sku": tesco_sku,
                    "wooper_sku": str(
                        catalogue_row.get("wooper_sku")
                        or source_row.get("wooper_sku")
                        or ""
                    ).strip().upper(),
                    "barcode": normalise_barcode(source_row.get("barcode")),
                    "title": str(source_row.get("title") or "").strip(),
                }
            )
        if not rows:
            raise ValueError(f"{event_name} has no Tesco SKUs to save.")
        prepared.append(
            {
                "event_name": event_name,
                "start_date": start_date,
                "end_date": end_date,
                "rows": rows,
            }
        )
        existing.add(identity)
    return [
        save_tesco_event(
            event,
            source={
                "external_nomination_file": str(source_filename or "").strip(),
                "imported_outside_tool": True,
            },
        )
        for event in prepared
    ]


def delete_tesco_event(event_id):
    if supabase_enabled():
        event_id = validate_worktable_id(event_id)
        rows = supabase_request(
            "DELETE",
            "promotion_tesco_events",
            params={"id": f"eq.{event_id}"},
            prefer="return=representation",
        )
        return rows[0] if rows else None
    events = tesco_events()
    remaining = [event for event in events if str(event.get("id")) != event_id]
    if len(remaining) == len(events):
        return None
    removed = next(event for event in events if str(event.get("id")) == event_id)
    write_json_file(TESCO_EVENTS_PATH, {"events": remaining})
    nomination_path = TESCO_NOMINATIONS_DIR / f"{event_id}.xlsx"
    if nomination_path.exists():
        nomination_path.unlink()
    return removed


def valid_basic_authorization(header_value):
    if not header_value or not header_value.startswith("Basic "):
        return False
    try:
        decoded = base64.b64decode(
            header_value.split(" ", 1)[1],
            validate=True,
        ).decode("utf-8")
    except (ValueError, UnicodeDecodeError):
        return False
    username, separator, password = decoded.partition(":")
    if not separator:
        return False
    return (
        hmac.compare_digest(username, PROMOTION_AUTH_USERNAME)
        and hmac.compare_digest(password, PROMOTION_AUTH_PASSWORD)
    )


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(STATIC_DIR), **kwargs)

    def send_json(self, payload, status=200):
        body = json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_bytes(self, body, content_type, filename, status=200, extra_headers=None):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        for name, value in (extra_headers or {}).items():
            self.send_header(name, str(value))
        self.end_headers()
        self.wfile.write(body)

    def read_json(self):
        length = int(self.headers.get("Content-Length", "0"))
        return json.loads(self.rfile.read(length) or b"{}")

    def require_authorization(self, path):
        if path == "/api/health" or not PROMOTION_REQUIRE_AUTH:
            return True
        if valid_basic_authorization(self.headers.get("Authorization")):
            return True
        self.send_response(401)
        self.send_header("WWW-Authenticate", 'Basic realm="Promotion Tool"')
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", "0")
        self.end_headers()
        return False

    def do_GET(self):
        parsed = urlparse(self.path)
        if not self.require_authorization(parsed.path):
            return
        if parsed.path == "/api/health":
            self.send_json(
                {
                    "ok": True,
                    "mode": "promotion-tool",
                    "dataSource": PROMOTION_DATA_SOURCE,
                    "supabase": supabase_enabled(),
                    "authRequired": PROMOTION_REQUIRE_AUTH,
                }
            )
            return
        if parsed.path == "/api/sample":
            payload = json.loads(SAMPLE_PATH.read_text(encoding="utf-8"))
            inventory = wooper_inventory()
            exact_ca_prices, _ = channeladvisor_prices()
            _, aligned_ca_prices = align_channeladvisor_prices(
                inventory,
                mappings=persisted_mapping_dictionary("channeladvisor"),
                exact_prices=exact_ca_prices,
            )
            for row in payload.get("rows", []):
                sku = str(row.get("sku") or "").strip().upper()
                performance = inventory.get(sku, {})
                for field in (
                    "sold_qty",
                    "sales_amt",
                    "return_rate",
                    "normal_margin",
                    "lifetime_profit_margin",
                ):
                    if performance.get(field) is not None:
                        row[field] = performance[field]
                row["offer_price"] = row.get("price")
                row["ca_price"] = (
                    exact_ca_prices.get(str(row.get("platform_sku") or "").upper())
                    or aligned_ca_prices.get(str(row.get("sku") or "").upper())
                    or row.get("price")
                )
            self.send_json(payload)
            return
        if parsed.path == "/api/config":
            inventory = wooper_inventory()
            settings = platform_settings()
            exact_ca_prices, _ = channeladvisor_prices()
            ca_mappings = persisted_mapping_dictionary("channeladvisor")
            ca_alignment, _ = align_channeladvisor_prices(
                inventory,
                mappings=ca_mappings,
                exact_prices=exact_ca_prices,
            )
            ca_alignment_summary = {
                status: sum(
                    row["mapping_status"] == status
                    for row in ca_alignment.values()
                )
                for status in ("mapped", "unresolved", "non_existing")
            }
            protection_rows = active_protection_rows()
            self.send_json(
                {
                    "defaultCommissions": {
                        row["platform"]: row["default_commission"]
                        for row in settings
                    },
                    "platformSettings": settings,
                    "variableCommissionPlatforms": sorted(
                        VARIABLE_COMMISSION_PLATFORMS
                    ),
                    "wooperSkus": supabase_inventory_skus() or sorted(inventory),
                    "caPriceCount": len(exact_ca_prices),
                    "caAlignmentSummary": ca_alignment_summary,
                    "caUnresolved": unresolved_channeladvisor_rows(),
                    "mappingStorage": (
                        "supabase" if supabase_enabled() else "local"
                    ),
                    "promotionDataRefreshedAt": promotion_data_refreshed_at(),
                    "activeProtectionCount": len(protection_rows),
                }
            )
            return
        if parsed.path == "/api/platform-settings":
            try:
                self.send_json({"settings": platform_settings()})
            except RuntimeError as exc:
                self.send_json({"error": str(exc)}, status=503)
            return
        if parsed.path == "/api/tesco/status":
            offers = latest_tesco_offers()
            self.send_json(
                {
                    "offerCount": len(offers),
                    "offersCapturedAt": latest_tesco_offers_captured_at(offers),
                    "events": tesco_events(),
                    "level1": TESCO_LEVEL_1,
                    "storage": "supabase" if supabase_enabled() else "local",
                }
            )
            return
        if parsed.path == "/api/worktables":
            try:
                query = parse_qs(parsed.query)
                rows = list_saved_worktables(
                    platform=(query.get("platform") or [""])[0],
                    event_name=(query.get("event_name") or [""])[0],
                    created_on=(query.get("created_on") or [""])[0],
                    limit=(query.get("limit") or [100])[0],
                )
                self.send_json({"worktables": rows})
            except ValueError as exc:
                self.send_json({"error": str(exc)}, status=400)
            except RuntimeError as exc:
                self.send_json({"error": str(exc)}, status=503)
            return
        if parsed.path.startswith("/api/worktables/"):
            try:
                record = get_saved_worktable(parsed.path.rsplit("/", 1)[-1])
                if record is None:
                    self.send_json({"error": "Saved worktable not found."}, status=404)
                else:
                    self.send_json({"worktable": record})
            except ValueError as exc:
                self.send_json({"error": str(exc)}, status=400)
            except RuntimeError as exc:
                self.send_json({"error": str(exc)}, status=503)
            return
        super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        if not self.require_authorization(parsed.path):
            return
        try:
            if parsed.path == "/api/platform-settings":
                settings = save_platform_settings(self.read_json())
                self.send_json({"settings": settings})
                return
            if parsed.path == "/api/worktables":
                record = create_saved_worktable(self.read_json())
                self.send_json({"worktable": record}, status=201)
                return
            if parsed.path == "/api/calculate":
                payload = self.read_json()
                criteria = dict(payload.get("criteria", {}))
                if criteria.get("exclude_current_protection_list"):
                    criteria["protected_skus"] = sorted(
                        {
                            str(row.get("sku") or "").strip().upper()
                            for row in active_protection_rows()
                            if row.get("sku")
                        }
                    )
                candidates = [
                    calculate_candidate(row, criteria)
                    for row in payload.get("rows", [])
                ]
                self.send_json({"candidates": candidates})
                return
            if parsed.path == "/api/import":
                length = int(self.headers.get("Content-Length", "0"))
                data = self.rfile.read(length)
                filename = self.headers.get("X-Filename", "").lower()
                if filename.endswith(".csv"):
                    rows = rows_from_csv(data)
                elif filename.endswith((".xlsx", ".xlsm")):
                    rows = rows_from_workbook(data)
                else:
                    self.send_json(
                        {"error": "Use a CSV, XLSX, or XLSM file."}, status=400
                    )
                    return
                platform = self.headers.get("X-Platform", "")
                mapped = map_imported_rows(rows, platform=platform)
                if clean_header(platform) == "tesco":
                    capture_tesco_offers(rows, mapped)
                unresolved = [
                    row for row in mapped if row.get("mapping_status") == "unresolved"
                ]
                self.send_json(
                    {
                        "rows": mapped,
                        "sourceRows": len(rows),
                        "unresolved": unresolved,
                    }
                )
                return
            if parsed.path == "/api/tesco/import-candidates":
                length = int(self.headers.get("Content-Length", "0"))
                data = self.rfile.read(length)
                filename = self.headers.get("X-Filename", "").lower()
                if filename.endswith(".csv"):
                    rows = rows_from_csv(data)
                elif filename.endswith((".xlsx", ".xlsm")):
                    rows = rows_from_workbook(data)
                else:
                    self.send_json({"error": "Use the exported CSV, XLSX, or XLSM file."}, status=400)
                    return
                candidates = tesco_candidates_from_rows(rows, inventory=wooper_inventory())
                if not candidates:
                    self.send_json(
                        {"error": "No Tesco rows have Yes in the Approval column."},
                        status=400,
                    )
                    return
                self.send_json(
                    {
                        "candidates": candidates,
                        "approvedRows": len(candidates),
                        "sourceRows": len(rows),
                    }
                )
                return
            if parsed.path == "/api/tesco/import-catalogue":
                length = int(self.headers.get("Content-Length", "0"))
                data = self.rfile.read(length)
                filename = self.headers.get("X-Filename", "").lower()
                if not filename.endswith((".xlsx", ".xlsm")):
                    self.send_json({"error": "Use the latest Tesco Catalogue XLSX file."}, status=400)
                    return
                rows = tesco_catalogue_from_workbook(
                    data,
                    inventory=wooper_inventory(),
                    mappings=combined_platform_mappings("Tesco"),
                )
                if not rows:
                    self.send_json({"error": "No Tesco Catalogue rows were found."}, status=400)
                    return
                rows = capture_tesco_catalogue(rows)
                self.send_json({"catalogue": rows, "sourceRows": len(rows)})
                return
            if parsed.path == "/api/tesco/import-external":
                length = int(self.headers.get("Content-Length", "0"))
                data = self.rfile.read(length)
                filename = self.headers.get("X-Filename", "").lower()
                if not filename.endswith((".xlsx", ".xlsm")):
                    self.send_json({"error": "Use a Tesco nomination XLSX or XLSM file."}, status=400)
                    return
                catalogue = latest_tesco_catalogue()
                if not catalogue:
                    self.send_json(
                        {
                            "error": (
                                "Import the latest Tesco Catalogue once before "
                                "recording an external nomination form."
                            )
                        },
                        status=400,
                    )
                    return
                result = tesco_external_nomination_from_workbook(data, catalogue)
                self.send_json(result)
                return
            if parsed.path == "/api/tesco/events/import":
                payload = self.read_json()
                records = import_external_tesco_events(
                    payload,
                    source_filename=str(payload.get("source_filename") or ""),
                )
                self.send_json({"events": records})
                return
            if parsed.path == "/api/tesco/check-overlap":
                payload = self.read_json()
                self.send_json(
                    {
                        "conflicts": tesco_conflicts(
                            str(payload.get("start_date") or ""),
                            str(payload.get("end_date") or ""),
                        ),
                        "offers": latest_tesco_offers(),
                    }
                )
                return
            if parsed.path == "/api/tesco/generate":
                payload = self.read_json()
                workbook = build_tesco_nomination_workbook(payload)
                record = save_tesco_event(payload, workbook)
                slug = re.sub(r"[^a-z0-9]+", "-", record["event_name"].lower()).strip("-")
                self.send_bytes(
                    workbook,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    f"tesco-{slug or 'event'}-nomination.xlsx",
                    extra_headers={"X-Tesco-Event-Id": record["id"]},
                )
                return
            if parsed.path == "/api/import-commissions":
                length = int(self.headers.get("Content-Length", "0"))
                data = self.rfile.read(length)
                filename = self.headers.get("X-Filename", "").lower()
                if filename.endswith(".csv"):
                    rows, conflicts = normalise_commission_rows(
                        rows_from_csv(data)
                    )
                elif filename.endswith((".xlsx", ".xlsm")):
                    rows, conflicts = commission_rows_from_workbook(data)
                else:
                    self.send_json(
                        {"error": "Use a CSV, XLSX, or XLSM file."},
                        status=400,
                    )
                    return
                if not rows:
                    self.send_json(
                        {
                            "error": (
                                "No commission rows were found. Expected Platform, "
                                "Category, SKU, and Commission Rate columns."
                            )
                        },
                        status=400,
                    )
                    return
                if conflicts:
                    self.send_json(
                        {
                            "error": (
                                "Conflicting commission rates were found for the "
                                "same platform and SKU."
                            ),
                            "conflicts": conflicts,
                        },
                        status=400,
                    )
                    return
                platform_counts = {}
                for row in rows:
                    platform = row["platform"]
                    platform_counts[platform] = (
                        platform_counts.get(platform, 0) + 1
                    )
                self.send_json(
                    {
                        "rows": rows,
                        "platformCounts": platform_counts,
                    }
                )
                return
            if parsed.path == "/api/mappings":
                payload = self.read_json()
                mapping_scope = str(
                    payload.get("mapping_scope") or "platform"
                ).strip().lower()
                if mapping_scope not in {"channeladvisor", "platform"}:
                    self.send_json({"error": "Invalid mapping scope."}, status=400)
                    return
                platform = str(payload.get("platform") or "").strip()
                inventory_skus = {
                    str(sku).strip().upper()
                    for sku in (
                        supabase_inventory_skus() or wooper_inventory()
                    )
                    if str(sku).strip()
                }
                invalid_mappings = []
                valid_mappings = []
                for item in payload.get("mappings", []):
                    platform_sku = str(item.get("platform_sku") or "").strip().upper()
                    wooper_sku = str(item.get("wooper_sku") or "").strip().upper()
                    if platform_sku and item.get("non_existing"):
                        valid_mappings.append(
                            {
                                "external_sku": platform_sku,
                                "wooper_sku": None,
                                "status": "non_existing",
                            }
                        )
                        continue
                    if not platform_sku or not wooper_sku:
                        continue
                    if wooper_sku not in inventory_skus:
                        invalid_mappings.append(
                            {
                                "platform_sku": platform_sku,
                                "wooper_sku": wooper_sku,
                            }
                        )
                    else:
                        valid_mappings.append(
                            {
                                "external_sku": platform_sku,
                                "wooper_sku": wooper_sku,
                                "status": "mapped",
                            }
                        )
                save_persisted_mappings(
                    valid_mappings,
                    mapping_scope,
                    platform,
                )
                remapped = map_imported_rows(
                    payload.get("rows", []),
                    platform=platform,
                )
                self.send_json(
                    {
                        "rows": remapped,
                        "invalidMappings": invalid_mappings,
                        "savedCount": len(valid_mappings),
                        "unresolved": [
                            row
                            for row in remapped
                            if row.get("mapping_status") == "unresolved"
                        ],
                        "caUnresolved": unresolved_channeladvisor_rows(),
                    }
                )
                return
            self.send_json({"error": "Not found"}, status=404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, status=400)

    def do_PATCH(self):
        parsed = urlparse(self.path)
        if not self.require_authorization(parsed.path):
            return
        try:
            if parsed.path.startswith("/api/tesco/events/"):
                event_id = parsed.path.rsplit("/", 1)[-1]
                record = update_tesco_event(event_id, self.read_json())
                if record is None:
                    self.send_json({"error": "Saved Tesco event not found."}, status=404)
                else:
                    self.send_json({"event": record})
                return
            self.send_json({"error": "Not found"}, status=404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, status=400)

    def do_DELETE(self):
        parsed = urlparse(self.path)
        if not self.require_authorization(parsed.path):
            return
        if parsed.path.startswith("/api/tesco/events/"):
            event_id = parsed.path.rsplit("/", 1)[-1]
            removed = delete_tesco_event(event_id)
            if removed is None:
                self.send_json({"error": "Saved Tesco event not found."}, status=404)
            else:
                self.send_json({"removed": removed})
            return
        self.send_json({"error": "Not found"}, status=404)


def main():
    if PROMOTION_DATA_SOURCE == "supabase" and not supabase_enabled():
        raise RuntimeError(
            "Supabase data source requires SUPABASE_URL and "
            "SUPABASE_SERVICE_ROLE_KEY."
        )
    if PROMOTION_REQUIRE_AUTH and not (
        PROMOTION_AUTH_USERNAME and PROMOTION_AUTH_PASSWORD
    ):
        raise RuntimeError(
            "PROMOTION_REQUIRE_AUTH is enabled but the authentication "
            "username or password is missing."
        )
    port = int(os.environ.get("PORT", "8877"))
    host = os.environ.get("HOST", "127.0.0.1")
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"Promotion tool running at http://{host}:{port}")
    server.serve_forever()


if __name__ == "__main__":
    main()
