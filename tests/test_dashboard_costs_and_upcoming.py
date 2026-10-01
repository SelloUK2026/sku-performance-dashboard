from __future__ import annotations

import sys
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook


APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

import app  # noqa: E402
import import_to_supabase as importer  # noqa: E402


class DashboardCostsAndUpcomingTests(unittest.TestCase):
    def test_freight_export_uses_shared_final_values_and_sorts_skus(self):
        rows = app.freight_export_rows(
            {
                "inventory": {
                    "B-UK": {
                        "sku": "B-UK",
                        "merchant_shipping_cost": 2.43,
                        "suggested_freight": 2.43,
                    },
                    "A-UK": {
                        "sku": "A-UK",
                        "merchant_shipping_cost": 1.69,
                        "suggested_freight": 2.125,
                    },
                },
                "freight": {
                    "B-UK": {"sku": "B-UK", "avg_actual_freight": 5.28},
                    "A-UK": {"sku": "A-UK", "avg_actual_freight": 2.125},
                    "OLD-UK": {"sku": "OLD-UK", "avg_actual_freight": 9.99},
                },
            }
        )

        self.assertEqual([row["Wooper SKU"] for row in rows], ["A-UK", "B-UK"])
        self.assertEqual(rows[0]["Wooper Freight"], 1.69)
        self.assertEqual(rows[0]["Eligible Average Freight"], 2.125)
        self.assertEqual(rows[0]["Suggested Freight"], 2.125)
        self.assertEqual(rows[1]["Suggested Freight"], 2.43)

    def test_freight_export_workbook_includes_bilingual_guide(self):
        payload = app.freight_export_workbook(
            {
                "inventory": {
                    "A-UK": {
                        "sku": "A-UK",
                        "merchant_shipping_cost": 1.69,
                        "suggested_freight": 2.125,
                    }
                },
                "freight": {
                    "A-UK": {"sku": "A-UK", "avg_actual_freight": 2.125}
                },
            }
        )

        workbook = load_workbook(BytesIO(payload), data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Freight", "Guide 计算说明"])
        self.assertEqual(
            [cell.value for cell in workbook["Freight"][1]],
            [
                "Wooper SKU",
                "Wooper Freight",
                "Eligible Average Freight",
                "Suggested Freight",
            ],
        )
        self.assertEqual(list(workbook["Freight"].tables), [])
        self.assertEqual(workbook["Freight"].auto_filter.ref, "A1:D2")
        guide_text = " ".join(
            str(cell.value or "") for row in workbook["Guide 计算说明"] for cell in row
        )
        self.assertIn("Amazon(UK) FBA", guide_text)
        self.assertIn("greater than 5", guide_text)
        self.assertIn("大于 5", guide_text)
        self.assertIn("£1.69", guide_text)

    def test_cogs_percentage_uses_requested_sales_base(self):
        rows = [
            {
                "platform": "Test Platform",
                "sku_qty": 2,
                "sales_amt": 100,
                "cogs": 25,
                "extra_freight": 10,
                "promo_rebate": -5,
                "selling_fee": 0,
                "ads_fee": 0,
                "resend_amt": 0,
                "refund_amt": 0,
                "profit_incl_rn": 30,
            }
        ]

        result = app.aggregate_sales_rows(rows)

        platform_row = next(row for row in result if row["platform name"] == "Test Platform")
        self.assertAlmostEqual(platform_row["cogs_pct"], 25 / 105)

    def test_upcoming_stock_export_maps_required_fields(self):
        headers = [
            "Product SKU",
            "Container No.1 Stock Qty",
            "Container 1 ETA (WH)",
            "Container No.2 Stock Qty",
            "Container 2 ETA (WH)",
            "Reorder Placed Date",
            "Production Scheduled Quantity",
        ]
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        sheet.append(["bm2006-uk", 100, "2026-09-30", 400, "2026-10-07", "2026-08-11", 250])

        export_path = APP_DIR / "tests" / "fixtures" / "upcoming_test.xlsx"
        self.addCleanup(export_path.unlink, missing_ok=True)
        try:
            workbook.save(export_path)
            previous_path = importer.UPCOMING_STOCK_FILE
            importer.UPCOMING_STOCK_FILE = export_path
            try:
                rows = importer.build_upcoming_stock()
            finally:
                importer.UPCOMING_STOCK_FILE = previous_path
        finally:
            workbook.close()

        self.assertEqual(rows[0]["sku"], "BM2006-UK")
        self.assertEqual(rows[0]["container_1_stock_qty"], 100)
        self.assertEqual(rows[0]["container_1_eta"], "2026-09-30")
        self.assertEqual(rows[0]["container_2_stock_qty"], 400)
        self.assertEqual(rows[0]["container_2_eta"], "2026-10-07")
        self.assertEqual(rows[0]["reorder_placed_date"], "2026-08-11")
        self.assertEqual(rows[0]["production_scheduled_qty"], 250)


if __name__ == "__main__":
    unittest.main()

